import fitz  # PyMuPDF
import os
from datetime import datetime
from docx import Document
from openai import OpenAI

# =========================
# 🔑 AI配置（DeepSeek）
# =========================
API_KEY = os.environ.get("DEEPSEEK_API_KEY")


def get_client():
    if not API_KEY:
        raise RuntimeError("❌ 未检测到 DEEPSEEK_API_KEY 环境变量，请先设置 DeepSeek API Key")
    return OpenAI(api_key=API_KEY, base_url="https://api.deepseek.com")


if not API_KEY:
    print("错误：未检测到 DEEPSEEK_API_KEY 环境变量，请先设置 DeepSeek API Key")
    raise SystemExit(1)

# =========================
# 📄 读取PDF
# =========================
def read_pdf(file_path):
    doc = fitz.open(file_path)
    text = ""

    for page in doc:
        text += page.get_text()

    return text


# =========================
# 🤖 AI分析
# =========================
def ai_analyze(text):
    prompt = f"""
你是一个顶级科研助手，请分析以下论文内容：

请输出结构化内容：

1. 论文一句话总结
2. 研究背景
3. 研究方法
4. 核心结果
5. 创新点
6. 局限性
7. 应用方向
8. 通俗解释（给非专业人士）

论文内容：
{text}
"""

    response = get_client().chat.completions.create(
        model="deepseek-chat",
        messages=[
            {"role": "system", "content": "你是专业科研分析助手"},
            {"role": "user", "content": prompt}
        ],
        temperature=0.3
    )

    return response.choices[0].message.content


# =========================
# 📝 Word报告生成
# =========================
import re
from docx import Document

def save_docx(file_name, result):
    doc = Document()

    # 标题
    doc.add_heading("🧬 论文AI分析报告", level=1)

    doc.add_paragraph(f"文件：{file_name}")
    doc.add_paragraph("")

    lines = result.split("\n")

    for line in lines:

        line = line.strip()

        if not line:
            continue

        # ===== 处理标题：### =====
        if line.startswith("###"):
            title = line.replace("#", "").strip()
            doc.add_heading(title, level=2)

        # ===== 处理加粗：**text** =====
        elif "**" in line:
            clean_line = re.sub(r"\*\*(.*?)\*\*", r"\1", line)
            p = doc.add_paragraph()
            run = p.add_run(clean_line)
            run.bold = True

        # ===== 普通文本 =====
        else:
            doc.add_paragraph(line)

    output_file = file_name.replace(".pdf", "_AI报告.docx")
    doc.save(output_file)

    print("✅ 美化Word报告已生成:", output_file)


# =========================
# 🚀 主程序
# =========================
import glob

print("=" * 60)
print("🧬 Paper Analyzer v4.0 - Word批处理版")
print("=" * 60)

pdf_files = glob.glob("paper/*.pdf")

if not pdf_files:
    print("❌ paper文件夹里没有PDF")
    exit()

for pdf_file in pdf_files:
    print(f"\n📄 正在处理: {pdf_file}")

    text = read_pdf(pdf_file)

    print("🤖 正在调用AI分析...")
    result = ai_analyze(text)

    print("📝 正在生成Word报告...")
    save_docx(pdf_file, result)

print("\n🎉 全部处理完成！")
