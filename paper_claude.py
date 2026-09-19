"""
======================================================================
 🧬 Paper Analyzer Pro v7.0 — 网页版
======================================================================
 用法：
    python paper_claude.py
    
 然后浏览器打开 http://localhost:7860
======================================================================
"""

import fitz
import json
import re
import os
import sqlite3
import time
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import gradio as gr
from paper_pipeline import PipelineConfig, request_budget_summary, run_batch
from providers import DEFAULT_OUTPUT_TOKEN_BUDGETS, ProviderError, create_provider, provider_defaults, provider_names, sanitize_error
from exporters import export_csv, export_excel, export_individual_reports, export_json, export_word
from bilingual import translate_paper_for_display, translated_paper_text
from local_extractor import extract_local_paper
from local_summary import LANGUAGE_OPTIONS, compare_papers, extractive_summary, offline_language_notice
from local_search import LocalSearchIndex, SearchContext
from passage_retrieval import build_evidence_passages
from batch_manager import run_local_batch, update_batch_manifest
from report_modes import NO_WORD_MODE, REPORT_MODE_OPTIONS, report_layout_plan, select_word_papers
from grounded_qa import build_evidence_pack, render_evidence_pack
from retrieval import (
    HybridRetriever,
    HybridSearchResponse,
    LocalEmbeddingSemanticRetriever,
    LocalModelPathError,
    OptionalSemanticDependencyError,
    build_semantic_index,
)
from retrieval.lexical import LexicalRetriever
from retrieval.ui_runtime import (
    LOCAL_MODEL_CACHE,
    PROFILE_OPTIONS,
    default_semantic_ui_state,
)

load_dotenv(os.path.join(os.path.dirname(__file__), ".env"), override=False)

# ╔═══════════════════════════════════════════╗
# ║  配置区域                                  ║
# ╚═══════════════════════════════════════════╝

# ============================================================
#  核心功能函数（复用原有逻辑）
# ============================================================

def get_client(api_key, provider_name="DeepSeek", model=None, base_url=None):
    """兼容旧入口，但实际返回统一 LLMProvider。"""
    return create_provider(provider_name, model=model, base_url=base_url, api_key=api_key)


def read_pdf(file_path):
    try:
        doc = fitz.open(file_path)
    except Exception as e:
        raise Exception(f"无法打开 PDF: {e}")
    if doc.is_encrypted:
        doc.close()
        raise Exception("PDF 已加密，无法读取")
    text = ""
    for i in range(doc.page_count):
        page_text = doc[i].get_text()
        if page_text.strip():
            text += page_text + "\n"
    doc.close()
    if not text.strip():
        raise Exception("PDF 内容为空（可能是扫描版图片）")
    return text


def call_ai(client, messages, temperature=0.3, max_retries=2, max_output_tokens=None):
    output_limit = (
        DEFAULT_OUTPUT_TOKEN_BUDGETS.final_max_output_tokens
        if max_output_tokens is None else max_output_tokens
    )
    last_error = None
    for attempt in range(max_retries + 1):
        try:
            return client.generate(messages, temperature=temperature, timeout=60, max_output_tokens=output_limit)
        except Exception as e:
            last_error = sanitize_error(e, getattr(getattr(client, "config", None), "api_key", None))
            if attempt < max_retries:
                time.sleep((attempt + 1) * 3)
            else:
                raise Exception(f"AI 调用失败（已重试{max_retries}次）: {last_error}")


def analyze_single_paper(client, text, paper_name=""):
    truncated = text
    prompt = f"""请以专业科研分析助手的身份，对以下论文进行结构化分析。

论文名称：{paper_name}

请严格按照以下格式输出（** 表示加粗内容）：

**研究背景**
（2-3句话说明研究背景和问题）

**研究方法**
（说明采用的核心方法和技术路线）

**创新点**
（列出2-3个关键创新点）

**结论**
（总结核心发现和结论）

——————————
论文内容：
{truncated}
"""
    messages = [
        {"role": "system", "content": "你是专业的生物医学文献分析助手。请严格使用 **加粗** 标记输出结构标题，输出内容要专业、简洁、有深度。"},
        {"role": "user", "content": prompt}
    ]
    return call_ai(client, messages, temperature=0.3)


def generate_review(client, all_analyses):
    prompt = f"""你是一位顶级的科研综述专家。请根据以下多篇论文的分析结果，撰写一篇高质量的文献综述。

要求：
1. **研究背景** — 该领域的整体研究背景和意义
2. **方法分类** — 各篇论文使用的方法归纳分类
3. **对比分析** — 不同研究的异同点、优劣势对比
4. **研究不足** — 现有研究的共同局限
5. **未来方向** — 该领域未来的研究方向

写作规范：
- 用 **加粗** 标记以上5个标题
- 正文中关键术语适当用 **加粗** 强调
- 语言专业流畅，符合学术综述风格

各篇论文分析：
{all_analyses}
"""
    messages = [
        {"role": "system", "content": "你是顶级的科研综述写作专家。请严格使用 **加粗** 标记小标题和关键术语。"},
        {"role": "user", "content": prompt}
    ]
    return call_ai(client, messages, temperature=0.4)


def add_formatted_paragraph(doc, line):
    heading_keywords = [
        "研究背景", "背景", "方法", "方法分类", "创新点", "创新",
        "结论", "总结", "对比分析", "研究不足", "未来方向",
        "研究方法", "核心结果"
    ]
    clean = line.replace("###", "").replace("##", "").replace("#", "").strip()
    is_heading = False
    for kw in heading_keywords:
        if clean.startswith(kw) or clean == kw:
            is_heading = True
            break
        if clean == f"**{kw}**" or clean.startswith(f"**{kw}**"):
            is_heading = True
            break
        no_num = re.sub(r"^\d+[\.\)、\s]*", "", clean)
        if no_num == f"**{kw}**" or no_num.startswith(f"**{kw}**") or no_num == kw:
            is_heading = True
            break
    if is_heading:
        heading_text = re.sub(r"\*\*(.*?)\*\*", r"\1", clean)
        heading_text = re.sub(r"^\d+[\.\)、\s]*", "", heading_text).strip()
        doc.add_heading(heading_text, level=2)
        return
    if "**" in line:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        parts = re.split(r"(\*\*[^*]+\*\*)", line)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                bold_text = part[2:-2]
                if bold_text.strip():
                    run = p.add_run(bold_text)
                    run.bold = True
                    run.font.color.rgb = RGBColor(0, 51, 153)
                    run.font.name = 'Times New Roman'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
            else:
                if part.strip():
                    run = p.add_run(part)
                    run.font.name = 'Times New Roman'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    else:
        clean_line = re.sub(r"^\d+[\.\)、\s]*", "", line).strip()
        if clean_line:
            p = doc.add_paragraph(clean_line)
            p.paragraph_format.space_after = Pt(2)


def save_docx(title, content, output_path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(2)
    h = doc.add_heading(title, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    time_p = doc.add_paragraph()
    time_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = time_p.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)
    doc.add_paragraph("")
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("———") or line.startswith("==="):
            continue
        add_formatted_paragraph(doc, line)
    doc.save(output_path)
    return output_path


# ============================================================
#  Gradio 处理函数
# ============================================================

def process_local_papers(
    pdf_files,
    progress=gr.Progress(),
    result_language="中文摘要＋英文证据（推荐）",
    report_mode="自动选择",
    resume_dir=None,
    pause_after=None,
    return_search_context=False,
):
    """本地离线解析：不创建 DeepSeek 客户端、不访问网络。"""
    if not pdf_files:
        result = ("❌ 请上传至少一个 PDF 文件", None)
        return (*result, None) if return_search_context else result
    paths = [pdf_file.name if hasattr(pdf_file, "name") else pdf_file for pdf_file in pdf_files]
    def report_progress(value, description):
        try:
            progress(value, desc=description)
        except TypeError:
            progress(value, description)

    try:
        batch = run_local_batch(
            paths,
            output_root="output",
            report_mode=report_mode,
            resume_dir=resume_dir,
            pause_after=pause_after,
            progress=report_progress,
            extractor=extract_local_paper,
        )
    except (OSError, ValueError) as exc:
        result = (f"❌ 本地批次处理失败：{exc}", None)
        return (*result, None) if return_search_context else result
    task_dir = batch.task_dir
    papers = batch.papers
    local_db_path = task_dir / "local_index.sqlite"
    indexed_sources = []
    with LocalSearchIndex(local_db_path) as local_index:
        for paper in papers:
            pages = paper.get("pages", []) or []
            if not pages or paper.get("errors"):
                continue
            local_index.index_pages(pages)
            local_index.index_passages(build_evidence_passages(paper))
            first_page = pages[0]
            source_file = first_page.source_file if hasattr(first_page, "source_file") else first_page.get("source_file")
            if source_file:
                indexed_sources.append(str(source_file))
    for paper in papers:
        summary = extractive_summary(paper)
        paper["extractive_summary"] = summary
    selection_records = select_word_papers(papers, max_papers=10)
    selection_updates = {}
    for paper, selection in zip(papers, selection_records):
        paper.update(selection)
        if paper.get("document_id"):
            selection_updates[paper["document_id"]] = {
                key: selection[key]
                for key in ("selection_strategy", "included_in_word", "selection_rank", "selection_score", "selection_reason")
            }
    batch.manifest = update_batch_manifest(task_dir, document_updates=selection_updates)
    comparison = compare_papers(papers)
    layout_plan = report_layout_plan(report_mode, len(papers))
    export_json(papers, task_dir / "local_papers.json")
    comparison_payload = {
        "documents": comparison,
        "inputs": batch.manifest.get("inputs", []),
        "selection_strategy": "metadata_completeness_verified_evidence_section_coverage_quality_document_id",
    }
    export_json(comparison_payload, task_dir / "comparison.json")
    if report_mode != NO_WORD_MODE:
        export_csv(comparison, task_dir / "comparison.csv")
    output_path = None
    if layout_plan["word_enabled"]:
        output_path = export_word(
            comparison,
            task_dir / "literature_analysis.docx",
            papers=papers,
            report_mode="Local Offline",
            report_layout_mode=report_mode,
            full_data_filename="comparison.xlsx / local_papers.json",
        )
    if layout_plan["individual_reports"]:
        individual_paths = export_individual_reports(comparison, papers, task_dir / "individual_reports")
        report_paths = {
            paper.get("document_id"): paper["individual_report"]
            for paper in papers
            if paper.get("document_id") and paper.get("individual_report")
        }
        batch.manifest = update_batch_manifest(task_dir, report_paths=report_paths)
    export_excel(
        comparison,
        task_dir / "comparison.xlsx",
        inputs=batch.manifest.get("inputs", []),
        documents=batch.manifest.get("documents", []),
    )
    lines = [
        "模式：Local Offline",
        f"结果语言：{result_language}",
        f"报告模式：{layout_plan['resolved_mode']}（布局：{layout_plan['layout']}）",
        offline_language_notice(result_language),
        f"本地结果目录：{task_dir}",
    ]
    if not layout_plan["word_enabled"]:
        lines.append("已按选择跳过 Word，仅导出 Excel/JSON。")
    if layout_plan["full_data_required"]:
        lines.append("完整批量数据：comparison.xlsx / local_papers.json；详细单篇报告：individual_reports/")
    if batch.manifest.get("status") == "paused":
        lines.append("⏸️ 批次已暂停；可使用同一结果目录和原始文件调用 resume_local_batch 继续。")
    lines.append("提示：本地规则提取暂不能稳定区分正文、表格标题和表格内容，重要结果请根据 PDF 页码回查原文。摘取式摘要仅选取原句，不代表 AI 生成或事实核验。")
    for paper in papers:
        lines.append(f"\n📄 {paper['file_name']}：{len(paper.get('pages', []))} 页，{len(paper.get('facts', []))} 条规则证据")
        lines.append(f"标题：{paper.get('title_candidate') or '未提取到'}；作者：{paper.get('author_candidate') or '未提取到'}；年份：{paper.get('year_candidate') or '未提取到'}")
        lines.append(f"研究对象：{paper.get('research_object') or '未提取到'}；样本量：{paper.get('sample_size') or '未提取到'}")
        methods = compare_papers([paper])[0].get("research_methods_keywords", [])
        method_labels = {"research_method": "研究方法", "model_fit": "模型拟合", "r_square": "模型拟合", "immunocapture_lc_ms": "免疫捕获 LC-MS", "affinity_purification_lc_ms": "亲和纯化 LC-MS", "single_dose_pk": "单次给药 PK", "multiple_dose_pk": "多次给药 PK", "randomized_controlled_trial": "随机对照试验", "dietary_intervention": "饮食干预", "randomization": "随机化", "double_blind": "双盲", "control_group": "对照组", "group_count": "分组数量"}
        lines.append(f"研究方法：{'; '.join(method_labels.get(method, method) for method in methods) or '未提取到'}")
        if paper.get("errors"):
            lines.extend(f"错误：{error}" for error in paper["errors"])
        if paper.get("warnings"):
            lines.extend(f"警告：{warning}" for warning in paper["warnings"])
        summary = paper.get("extractive_summary") or extractive_summary(paper)
        if summary["sentences"]:
            lines.append("摘取式摘要：")
            for item in summary.get("evidence", []):
                status = "已验证" if item.get("verified") else "未验证，请回查原文"
                pages = "、".join(str(page) for page in item.get("pdf_pages", [])) or "未知"
                lines.append(f"- [{item.get('section', '未标明章节')} · PDF 第 {pages} 页 · {status}] {item.get('text', '')}")
    result = ("\n".join(lines), str(output_path) if output_path else None)
    if not return_search_context:
        return result
    context = SearchContext(
        batch_id=task_dir.name,
        task_dir=str(task_dir),
        db_path=str(local_db_path),
        document_sources=tuple(dict.fromkeys(indexed_sources)),
        document_count=len(dict.fromkeys(indexed_sources)),
    ).to_dict()
    return (*result, context)


def _semantic_build_status_text(state: dict, message: str | None = None) -> str:
    if message:
        return message
    if not state.get("ready"):
        return "Hybrid Local 尚未就绪；当前批次仍可使用 Lexical / FTS5。"
    fingerprint = str(state.get("model_fingerprint") or "")[:12]
    return (
        "✅ Hybrid Local 已就绪（仅当前批次）\n\n"
        f"Profile：{state.get('model_profile') or '未指定'}；"
        f"维度：{state.get('embedding_dimension') or '未知'}；"
        f"段落：{state.get('indexed_passages', 0)}；"
        f"本次编码：{state.get('encoded', 0)}；缓存命中：{state.get('cached', 0)}；"
        f"模型指纹：{fingerprint or '未知'}"
    )


def _invalid_semantic_state(message: str | None = None) -> tuple[dict, str]:
    state = default_semantic_ui_state()
    state["status"] = "NOT_READY"
    state["warning"] = message
    return state, message or _semantic_build_status_text(state)


def reset_semantic_ui_state(reason: str = "") -> tuple[dict, str]:
    """重置会话级语义状态；不加载模型，也不读取论文内容。"""
    state = default_semantic_ui_state()
    state["warning"] = reason or None
    return state, reason or _semantic_build_status_text(state)


def build_semantic_index_for_ui(search_context, model_profile, model_dir):
    """由用户明确点击后，加载本地模型并为当前批次建立 semantic cache。"""
    context = SearchContext.from_value(search_context)
    if context is None:
        return _invalid_semantic_state("请先完成一次 Local Offline 文献处理，再构建当前批次语义索引。")
    if model_profile not in PROFILE_OPTIONS:
        return _invalid_semantic_state("请选择有效的本地模型 Profile。")
    if not isinstance(model_dir, str) or not model_dir.strip():
        return _invalid_semantic_state("请输入已存在的本地模型目录；不会自动下载模型。")
    try:
        model_path = Path(model_dir.strip()).expanduser()
        model_path_valid = model_path.is_dir()
    except (OSError, TypeError, ValueError):
        model_path_valid = False
        model_path = None
    if not model_path_valid:
        return _invalid_semantic_state("本地模型目录无效；请确认目录已存在且包含离线模型文件。")
    try:
        encoder, cache_hit = LOCAL_MODEL_CACHE.get_or_load(model_profile, model_path)
        stats = build_semantic_index(
            context,
            encoder,
            source_files=context.document_sources,
        )
    except OptionalSemanticDependencyError:
        return _invalid_semantic_state(
            "Hybrid Local 需要可选 semantic 依赖；Lexical / FTS5 仍可正常使用。可参考 requirements-semantic.txt。"
        )
    except LocalModelPathError:
        return _invalid_semantic_state("本地模型目录无效；不会根据模型名称访问网络。")
    except (OSError, ValueError, RuntimeError):
        return _invalid_semantic_state("本地语义索引构建失败；Lexical / FTS5 仍可使用。")
    except Exception:
        # UI 层不泄露模型路径、堆栈或底层依赖细节。
        return _invalid_semantic_state("本地语义索引构建失败；Lexical / FTS5 仍可使用。")

    state = default_semantic_ui_state()
    state.update(
        {
            "enabled": True,
            "ready": not bool(stats.get("failed")) and int(stats.get("total_passages", 0)) > 0,
            "status": "READY_FOR_CURRENT_BATCH" if not stats.get("failed") else "PARTIAL",
            "model_profile": model_profile,
            "model_fingerprint": str(stats.get("model_fingerprint") or encoder.model_fingerprint),
            "batch_id": context.batch_id,
            "embedding_dimension": int(stats.get("dimension", getattr(encoder, "dimension", 0)) or 0),
            "indexed_passages": int(stats.get("total_passages", 0)),
            "encoded": int(stats.get("encoded", 0)),
            "cached": int(stats.get("cached", 0)),
            "warning": "；".join(stats.get("warnings") or []) or None,
        }
    )
    if stats.get("failed"):
        state["ready"] = False
        state["warning"] = "部分 embedding 未能建立，已保持 Lexical / FTS5 可用。"
    elif not state["indexed_passages"]:
        state["ready"] = False
        state["status"] = "EMPTY"
        state["warning"] = "当前批次没有可建立语义索引的证据段落；Lexical / FTS5 仍可使用。"
    return state, _semantic_build_status_text(state)


def _format_retrieval_response(retrieval_response, *, warning: str | None = None) -> str:
    rows = [result.to_dict() for result in retrieval_response.results]
    if not rows:
        return warning or "未找到匹配内容。"
    formatted = []
    for row in rows:
        start = row["pdf_page_start"]
        end = row["pdf_page_end"]
        page_label = f"PDF 第 {start} 页" if start == end else f"PDF 第 {start}–{end} 页"
        formatted.append(
            f"[{row['rank']}] {row['section']} · {page_label}\n"
            f"{row['source_file']}\n{row['text']}"
        )
    prefix = warning + "\n\n" if warning else ""
    return prefix + "\n\n---\n\n".join(formatted)


def retrieve_local_response(
    query,
    search_context=None,
    limit=20,
    retrieval_mode="Lexical / FTS5",
    semantic_ui_state=None,
):
    """唯一的当前批次 retrieval 入口；Q&A 和搜索 UI 共用它。"""
    if not query or not query.strip():
        return None, "请输入关键词或精确短语。"
    context = SearchContext.from_value(search_context)
    if context is None:
        return None, "请先完成一次 Local Offline 文献处理。"
    task_dir = Path(context.task_dir)
    db_path = Path(context.db_path)
    try:
        task_dir_resolved = task_dir.resolve()
        db_path_resolved = db_path.resolve()
    except OSError:
        return None, "❌ 当前批次索引不存在或任务目录已被删除，无法搜索。"
    if (
        not task_dir_resolved.is_dir()
        or db_path_resolved.name != "local_index.sqlite"
        or db_path_resolved.parent != task_dir_resolved
        or not db_path_resolved.is_file()
    ):
        return None, "❌ 当前批次索引不存在或任务目录已被删除，无法搜索。"
    try:
        with LocalSearchIndex(db_path_resolved) as local_index:
            if local_index.passage_count() == 0:
                return (
                    HybridSearchResponse(
                        results=[],
                        mode="lexical",
                        semantic_status="unavailable",
                        warnings=["当前批次没有可用的证据段落索引。"],
                    ),
                    None,
                )
            lexical_retriever = LexicalRetriever(local_index)
            if retrieval_mode != "Hybrid Local":
                retrieval_response = HybridRetriever(lexical_retriever).search(
                    query,
                    limit=limit,
                    source_files=context.document_sources,
                )
                return retrieval_response, None

            state = semantic_ui_state if isinstance(semantic_ui_state, dict) else default_semantic_ui_state()
            if not state.get("ready") or state.get("batch_id") != context.batch_id:
                retrieval_response = HybridRetriever(lexical_retriever).search(
                    query,
                    limit=limit,
                    source_files=context.document_sources,
                )
                retrieval_response.warnings.append("Hybrid Local 尚未针对当前批次就绪，已回退为 Lexical / FTS5。")
                return retrieval_response, None
            encoder = LOCAL_MODEL_CACHE.get(state.get("model_fingerprint"))
            if encoder is None:
                retrieval_response = HybridRetriever(lexical_retriever).search(
                    query,
                    limit=limit,
                    source_files=context.document_sources,
                )
                retrieval_response.warnings.append("当前语义模型缓存已失效，已回退为 Lexical / FTS5。")
                return retrieval_response, None
            passages = local_index.list_passages(source_files=context.document_sources)
            retrieval_response = HybridRetriever(
                lexical_retriever,
                LocalEmbeddingSemanticRetriever(context, encoder),
                allowed_passage_ids={item["passage_id"] for item in passages},
            ).search(
                query,
                limit=limit,
                source_files=context.document_sources,
            )
            if retrieval_response.mode != "hybrid":
                retrieval_response.warnings.append("本地语义检索暂不可用，已回退为 Lexical / FTS5。")
            return retrieval_response, None
    except (OSError, ValueError, sqlite3.Error):
        return None, "❌ 当前批次索引无法读取，请重新完成 Local Offline 文献处理。"


def search_local_index(
    query,
    search_context=None,
    limit=20,
    retrieval_mode="Lexical / FTS5",
    semantic_ui_state=None,
):
    """只查询当前批次 SQLite 索引；不访问网络或历史全局索引。"""
    retrieval_response, error = retrieve_local_response(
        query,
        search_context=search_context,
        limit=limit,
        retrieval_mode=retrieval_mode,
        semantic_ui_state=semantic_ui_state,
    )
    if error:
        return error
    warning = None
    if retrieval_response.mode == "hybrid":
        warning = "检索模式：Hybrid Local / RRF"
    elif retrieval_mode == "Hybrid Local" and retrieval_response.warnings:
        warning = "⚠️ " + retrieval_response.warnings[-1]
    return _format_retrieval_response(retrieval_response, warning=warning)


def search_local_index_for_ui(query, retrieval_mode, search_context, semantic_ui_state):
    return search_local_index(
        query,
        search_context=search_context,
        retrieval_mode=retrieval_mode,
        semantic_ui_state=semantic_ui_state,
    )


def answer_evidence_only_for_ui(question, retrieval_mode, search_context, semantic_ui_state):
    """Evidence Only UI：只检索并回填原文，不创建或调用 Provider。"""
    retrieval_response, error = retrieve_local_response(
        question,
        search_context=search_context,
        limit=8,
        retrieval_mode=retrieval_mode,
        semantic_ui_state=semantic_ui_state,
    )
    if error:
        return error
    pack = build_evidence_pack(question, retrieval_response, search_context=search_context)
    return render_evidence_pack(pack)


def _selected_pdf_paths(pdf_files, selected_files):
    """按界面勾选结果筛选文件名；不会读取文件内容。"""
    if not pdf_files or not selected_files:
        return []
    selected = {Path(str(value)).name for value in selected_files}
    paths = [pdf_file.name if hasattr(pdf_file, "name") else pdf_file for pdf_file in pdf_files]
    return [path for path in paths if Path(str(path)).name in selected]


def estimate_ai_plan(pdf_files, selected_files, provider_name="DeepSeek", model=""):
    """本地估算 AI 请求规模；不创建 Provider、不发送网络请求。"""
    paths = _selected_pdf_paths(pdf_files, selected_files)
    if not paths:
        return "请选择需要 AI 分析的具体论文；当前仅显示估算，不会创建 Provider。"
    config = PipelineConfig(max_retries=0, repair_attempts=0)
    chunk_count = 0
    output_tokens = 0
    details = []
    base_requests = 0
    try:
        for path in paths:
            paper = extract_local_paper(path, config=config)
            count = len(paper.get("chunks", []))
            chunk_count += count
            summary = request_budget_summary(count, config, include_batch_review=False)
            base_requests += summary["base_request_count"]
            output_tokens += summary["theoretical_max_output_tokens"]
            details.append(f"{Path(path).name}：{count} 个 chunk，{summary['base_request_count']} 次（分块+单篇归并）")
    except (OSError, ValueError) as exc:
        return f"❌ 无法完成本地请求估算：{exc}"
    final_requests = 1 if paths else 0
    total_requests = base_requests + final_requests
    output_tokens += config.final_max_output_tokens
    lines = [
        f"AI 预览（{provider_name} / {model or '未填写模型'}）：仅本地估算，尚未创建 Provider。",
        f"已选择 {len(paths)} 篇，预计 {chunk_count} 个 chunk；基础请求约 {total_requests} 次（含最终汇总）。",
        f"输出预算上限：约 {output_tokens} tokens；max_retries=0，repair_attempts=0；hard_request_limit 不会自动扩大。",
        *details,
    ]
    return "\n".join(lines)


def _ai_selection_choices(pdf_files):
    if not pdf_files:
        return gr.update(choices=[], value=[])
    paths = [pdf_file.name if hasattr(pdf_file, "name") else pdf_file for pdf_file in pdf_files]
    choices = list(dict.fromkeys(Path(str(path)).name for path in paths))
    return gr.update(choices=choices, value=[])


def process_papers(
    pdf_files,
    api_key,
    mode=None,
    progress=gr.Progress(),
    provider_name="DeepSeek",
    model="",
    base_url="",
    result_language="中文摘要＋英文证据（推荐）",
    report_mode="自动选择",
    selected_ai_files=None,
    ai_confirmed=False,
):
    """主处理函数：接收上传的PDF，返回综述文本和Word文件"""

    if mode == "Local Offline":
        return process_local_papers(pdf_files, progress, result_language, report_mode=report_mode)

    try:
        defaults = provider_defaults(provider_name)
    except ProviderError as exc:
        return f"❌ {exc}", None
    resolved_key = api_key.strip() if api_key and api_key.strip() else os.environ.get(defaults.get("env_key")) if defaults.get("env_key") else None
    if provider_name != "Ollama" and not (resolved_key or "").strip():
        env_name = defaults.get("env_key")
        suffix = f"，或在环境变量中设置 {env_name}" if env_name else ""
        return f"❌ 请先填写 {provider_name} API Key{suffix}", None

    if not pdf_files:
        return "❌ 请上传至少一个 PDF 文件", None
    if not selected_ai_files:
        return "❌ AI 模式需先勾选具体论文；批量上传不会自动调用 AI。", None
    selected_paths = _selected_pdf_paths(pdf_files, selected_ai_files)
    if not selected_paths:
        return "❌ 未找到勾选的论文，请重新选择后再试。", None
    if not ai_confirmed:
        return "❌ 请先确认：选中的论文文本将发送至所选 AI Provider，可能产生请求或费用。", None

    # AI 前先完成同一批选中文件的本地提取；这一步不创建 Provider，也不联网。
    try:
        run_local_batch(selected_paths, output_root="output", report_mode="自动选择", progress=None)
    except (OSError, ValueError) as exc:
        return f"❌ AI 前置本地解析失败：{exc}", None

    try:
        client = get_client(api_key, provider_name=provider_name, model=model, base_url=base_url)
    except ProviderError as exc:
        if "缺少" in str(exc):
            env_name = defaults.get("env_key")
            suffix = f"，或在环境变量中设置 {env_name}" if env_name else ""
            return f"❌ 请先填写 {provider_name} API Key{suffix}", None
        return f"❌ {sanitize_error(exc, api_key)}", None

    log_lines = []
    paths = selected_paths
    progress(0.05, desc="正在进行分页提取与分段分析...")
    pipeline_config = PipelineConfig(failure_policy="fail_fast")
    pipeline_result = run_batch(
        client,
        paths,
        output_root="output",
        config=pipeline_config,
    )
    success_count = sum(not paper.get("errors") for paper in pipeline_result.papers)
    fail_count = len(pipeline_result.papers) - success_count
    for paper in pipeline_result.papers:
        log_lines.append(f"📄 {paper.get('file_name', '未知文件')}：{paper.get('chunk_count', 0)} 个分块")
    if result_language != "原文":
        translated = 0
        failed = 0
        for paper in pipeline_result.papers:
            if paper.get("analysis_status") == "complete":
                translation = translate_paper_for_display(client, paper, pipeline_config)
                translated += translation.get("translated_fields", 0) + translation.get("translated_evidence", 0)
                failed += int(translation.get("status") == "failed")
        if translated:
            log_lines.append(f"🌐 中文展示已生成：{translated} 项；英文证据和页码保持不变。")
        if failed:
            log_lines.append("⚠️ 部分翻译未通过数字/统计信息一致性检查，已回退显示英文原文。")
        export_json(pipeline_result.papers, pipeline_result.task_dir / "papers_bilingual.json")
    progress(0.85, desc="正在整理结构化文献综述...")
    review = _review_to_text(pipeline_result.final_review)
    translated_lines = [line for paper in pipeline_result.papers for line in translated_paper_text(paper)]
    if translated_lines:
        review += "\n\n**中文展示**\n" + "\n".join(f"- {line}" for line in translated_lines)
    log_lines.append(f"\n📝 综述处理完成（成功 {success_count} 篇，存在警告/错误 {fail_count} 篇）")
    log_lines.append(f"📁 中间结果：{pipeline_result.task_dir}")

    # 保存 Word
    progress(0.95, desc="正在生成 Word 文档...")
    try:
        if report_mode == NO_WORD_MODE:
            log_lines.append("⏭️ 已按报告模式跳过 Word，仅保留 JSON/Excel 等结构化结果。")
            progress(1.0, desc="完成！")
            return "\n".join(log_lines) + "\n\n" + review, None
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = os.path.join(str(pipeline_result.task_dir), f"literature_analysis_{timestamp}.docx")
        report_rows = compare_papers(pipeline_result.papers)
        export_word(report_rows, output_path, papers=pipeline_result.papers,
                    report_mode=f"AI Provider / {provider_name}", final_review=pipeline_result.final_review,
                    report_layout_mode=report_mode, full_data_filename="papers.json / review.json")
        log_lines.append(f"✅ Word 文档已生成")
    except Exception as e:
        return f"❌ Word 生成失败: {e}", None

    progress(1.0, desc="完成！")

    # 返回综述文本 + Word 文件路径
    result_text = "\n".join(log_lines) + "\n\n" + "="*50 + "\n\n" + review
    return result_text, output_path


def process_papers_for_ui(
    pdf_files,
    api_key,
    mode=None,
    progress=gr.Progress(),
    provider_name="DeepSeek",
    model="",
    base_url="",
    result_language="中文摘要＋英文证据（推荐）",
    report_mode="自动选择",
    selected_ai_files=None,
    ai_confirmed=False,
):
    """UI 适配层：额外返回当前批次搜索上下文，不改变旧处理函数返回值。"""
    if mode == "Local Offline":
        return process_local_papers(
            pdf_files,
            progress=progress,
            result_language=result_language,
            report_mode=report_mode,
            return_search_context=True,
        )
    result = process_papers(
        pdf_files,
        api_key,
        mode=mode,
        progress=progress,
        provider_name=provider_name,
        model=model,
        base_url=base_url,
        result_language=result_language,
        report_mode=report_mode,
        selected_ai_files=selected_ai_files,
        ai_confirmed=ai_confirmed,
    )
    return result[0], result[1], None


def process_papers_for_ui_with_semantic_state(
    pdf_files,
    api_key,
    mode=None,
    progress=gr.Progress(),
    provider_name="DeepSeek",
    model="",
    base_url="",
    result_language="中文摘要＋英文证据（推荐）",
    report_mode="自动选择",
    selected_ai_files=None,
    ai_confirmed=False,
):
    """新的 UI 入口：每次批次处理后明确清空旧 semantic ready 状态。"""
    result = process_papers_for_ui(
        pdf_files,
        api_key,
        mode=mode,
        progress=progress,
        provider_name=provider_name,
        model=model,
        base_url=base_url,
        result_language=result_language,
        report_mode=report_mode,
        selected_ai_files=selected_ai_files,
        ai_confirmed=ai_confirmed,
    )
    return result[0], result[1], result[2], default_semantic_ui_state()


def test_provider_connection(provider_name, model, base_url, api_key):
    """只发送最小健康检查，不发送论文文本。"""
    if provider_name == "Local Offline":
        return "Local Offline 不需要连接测试，不访问网络。"
    try:
        provider = create_provider(provider_name, model=model, base_url=base_url, api_key=api_key)
        provider.health_check()
        return f"✅ {provider_name} / {model} 连接测试成功（本次未发送论文文本）。"
    except ProviderError as exc:
        return f"❌ 连接测试失败：{sanitize_error(exc, api_key)}"


def _review_to_text(review):
    """将结构化综述转换为兼容现有文本框和 Word 导出的简洁文本。"""
    labels = {
        "research_theme_overview": "研究主题概述", "major_methods": "主要研究方法",
        "common_conclusions": "共同结论", "different_or_conflicting_conclusions": "不同或冲突结论",
        "research_innovations": "研究创新", "current_limitations": "现有局限",
        "research_gaps": "研究空白", "future_recommendations": "后续研究建议",
        "papers": "涉及的论文清单", "evidence": "可核查证据",
    }
    lines = []
    for key, label in labels.items():
        value = review.get(key)
        if value in (None, [], ""):
            continue
        lines.append(f"**{label}**")
        if isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    claim = item.get("claim") or item.get("text") or "未提取到"
                    page = item.get("pdf_page_start") or (item.get("pdf_pages") or [None])[0]
                    verified = "已验证" if item.get("verified") is True else "未验证，请回查原文"
                    lines.append(f"- {claim}（PDF 第 {page or '未知'} 页，{verified}）")
                else:
                    lines.append(f"- {item}")
        else:
            lines.append(str(value))
    if review.get("warnings"):
        lines.extend(["**警告**", *[f"- {item}" for item in review["warnings"]]])
    if review.get("errors"):
        lines.extend(["**错误摘要**", *[f"- {item}" for item in review["errors"]]])
    return "\n".join(lines) or "原文未明确说明。"


# ============================================================
#  Gradio 界面
# ============================================================

def build_ui():
    with gr.Blocks(
        title="🧬 科研文献智能分析系统",
        theme=gr.themes.Soft(
            primary_hue="blue",
            secondary_hue="slate",
        ),
        css="""
        .header { text-align: center; padding: 20px 0 10px 0; }
        .header h1 { font-size: 2em; color: #1e3a5f; }
        .header p { color: #666; margin-top: 5px; }
        .tip-box { background: #f0f7ff; border-left: 4px solid #3b82f6;
                   padding: 10px 15px; border-radius: 4px; margin: 10px 0; font-size: 0.9em; }
        footer { display: none !important; }
        """
    ) as demo:

        # 每个浏览器会话独立保存当前批次边界；不使用模块级全局索引。
        search_context_state = gr.State(None)
        # 仅保存小型、可序列化的状态；模型对象留在进程级有界 LRU 中。
        semantic_ui_state = gr.State(default_semantic_ui_state())

        # 顶部标题
        gr.HTML("""
        <div class="header">
            <h1>🧬 科研文献智能分析系统</h1>
            <p>上传 PDF 论文 → 本地解析或可选 AI 分析 → 结构化结果导出</p>
        </div>
        """)

        with gr.Row():
            # 左侧：输入区
            with gr.Column(scale=1):
                gr.Markdown("### ⚙️ 配置")

                mode_input = gr.Radio(
                    ["Local Offline", "AI Provider"], value="Local Offline",
                    label="运行模式", info="默认仅使用本地确定性功能，不访问网络"
                )

                result_language_input = gr.Radio(
                    list(LANGUAGE_OPTIONS), value="中文摘要＋英文证据（推荐）",
                    label="结果语言", info="Local Offline 不进行语义翻译，始终保留英文证据原文"
                )

                report_mode_input = gr.Dropdown(
                    list(REPORT_MODE_OPTIONS), value="自动选择", label="报告模式",
                    info="1篇不生成对比；4篇以上使用纵向文献库汇总；可选择不生成 Word",
                )

                provider_input = gr.Dropdown(
                    provider_names(), value="DeepSeek", label="AI Provider", visible=False
                )
                model_input = gr.Textbox(label="模型名称", value="deepseek-chat", visible=False)
                base_url_input = gr.Textbox(label="Base URL", value="https://api.deepseek.com", visible=False)
                api_key_input = gr.Textbox(
                    label="API Key（仅当前会话）",
                    placeholder="your_api_key_here",
                    type="password",
                    info="不会写入 .env 或结果文件",
                    visible=False,
                )
                connection_btn = gr.Button("🔌 测试连接（可能产生一次请求）", visible=False)
                connection_output = gr.Markdown(visible=False)

                gr.HTML('<div class="tip-box">💡 可选功能：论文文本将发送至 DeepSeek API。请勿上传涉密、敏感或未授权材料。<br>Local Offline 模式不需要 API Key，也不访问网络。</div>')

                gr.Markdown("### 📂 上传论文")

                pdf_input = gr.File(
                    label="选择 PDF 文件（可多选）",
                    file_count="multiple",
                    file_types=[".pdf"],
                )

                ai_selection_input = gr.CheckboxGroup(
                    choices=[], label="AI 分析论文（仅主动勾选的文件）", visible=False,
                    info="AI 模式会先完成本地提取；未勾选的论文不会发送给 Provider",
                )
                ai_confirm_input = gr.Checkbox(
                    label="我确认选中文本将发送至 AI Provider，可能产生请求或费用",
                    visible=False,
                )
                ai_plan_output = gr.Markdown(
                    "选择论文后显示本地请求估算。", visible=False,
                )

                gr.HTML('<div class="tip-box">📌 支持同时上传多篇论文；默认先离线解析，再按报告模式导出</div>')

                submit_btn = gr.Button(
                    "🚀 开始处理",
                    variant="primary",
                    size="lg"
                )

                gr.Markdown("### 📖 使用说明")
                gr.Markdown("""
1. 默认选择 Local Offline，无需 API Key
2. 上传一篇或多篇 PDF 论文
3. 点击「开始处理」查看本地解析和结构化对比
4. 如主动选择 AI Provider，再填写服务商、模型和配置

> ⚠️ 仅支持**文字版 PDF**，扫描版图片 PDF 无法读取
                """)

            # 右侧：输出区
            with gr.Column(scale=2):
                gr.Markdown("### 📊 分析结果")
                with gr.Tabs():
                    with gr.Tab("综合概览"):
                        result_text = gr.Textbox(
                            label="处理进度与结构化结果",
                            lines=25,
                            max_lines=40,
                            placeholder="点击「开始分析」后，这里会显示处理进度和结果概览...",
                        )
                    with gr.Tab("证据与页码"):
                        gr.Markdown("已验证：证据文本与 PDF 原文匹配；未验证：请根据 PDF 物理页码回查原文。source_type 仅表示程序对来源形态的保守判断。")
                    with gr.Tab("本地证据检索"):
                        gr.Markdown("本地证据检索：当前批次（仅搜索最近一次 Local Offline 处理结果，不检索历史任务）。\n\n默认使用 Lexical / FTS5；Hybrid Local 为可选的本地语义检索，不会自动加载或下载模型。")
                        retrieval_mode_input = gr.Radio(
                            ["Lexical / FTS5", "Hybrid Local"],
                            value="Lexical / FTS5",
                            label="检索模式",
                            info="选择 Hybrid Local 后仍需显式构建当前批次语义索引。",
                        )
                        with gr.Accordion(
                            "Hybrid Local 配置（仅主动构建时加载模型）",
                            visible=False,
                        ) as hybrid_settings:
                            model_profile_input = gr.Dropdown(
                                list(PROFILE_OPTIONS),
                                value="PubMedBERT",
                                label="模型 Profile",
                                info="仅接受已存在的本地模型目录；不会自动下载。",
                            )
                            model_dir_input = gr.Textbox(
                                label="本地模型目录",
                                placeholder="例如：D:\\models\\local-sentence-transformer",
                                info="仅当前会话使用，不会保存路径；请使用本机已有模型目录。",
                            )
                            build_semantic_btn = gr.Button("🧠 构建 / 加载当前批次语义索引")
                            semantic_status_output = gr.Markdown("Hybrid Local 尚未就绪；当前批次仍可使用 Lexical / FTS5。")
                        local_query = gr.Textbox(label="本地证据搜索", placeholder="关键词或精确短语（仅搜索本地索引）")
                        local_search_btn = gr.Button("🔎 搜索本地证据")
                        local_search_output = gr.Textbox(label="本地搜索结果", lines=8)
                        gr.Markdown("### 证据问答")
                        qa_question = gr.Textbox(
                            label="问题",
                            placeholder="例如：哪些证据提示延长热暴露会影响抗体稳定性？",
                        )
                        qa_answer_mode = gr.Radio(
                            ["Evidence Only"],
                            value="Evidence Only",
                            label="回答模式",
                            interactive=False,
                            info="只展示当前批次检索到的原文证据，不生成自然语言科研结论。",
                        )
                        qa_btn = gr.Button("📚 检索证据")
                        qa_output = gr.Textbox(label="Evidence Pack", lines=12)
                    with gr.Tab("导出结果"):
                        docx_output = gr.File(label="📥 下载 Word 文档", visible=True)
                        gr.Markdown("JSON、CSV、Excel 和 Word 会保存到本次任务的本地结果目录。")

        # 绑定事件
        submit_btn.click(
            fn=process_papers_for_ui_with_semantic_state,
            inputs=[pdf_input, api_key_input, mode_input, provider_input, model_input, base_url_input, result_language_input, report_mode_input, ai_selection_input, ai_confirm_input],
            outputs=[result_text, docx_output, search_context_state, semantic_ui_state],
        )

        pdf_input.change(
            _ai_selection_choices,
            inputs=[pdf_input],
            outputs=[ai_selection_input],
        )

        ai_selection_input.change(
            estimate_ai_plan,
            inputs=[pdf_input, ai_selection_input, provider_input, model_input],
            outputs=[ai_plan_output],
        )

        def update_ai_visibility(mode):
            visible = mode == "AI Provider"
            return [
                gr.update(visible=visible),
                gr.update(visible=visible),
                gr.update(visible=visible),
                gr.update(visible=visible, value=""),
                gr.update(visible=visible),
                gr.update(visible=visible),
                gr.update(visible=visible, value=[]),
                gr.update(visible=visible, value=False),
                gr.update(visible=visible, value="选择论文后显示本地请求估算。"),
            ]

        def update_provider_fields(name):
            defaults = provider_defaults(name)
            return defaults["model"], defaults["base_url"], ""

        mode_input.change(
            update_ai_visibility,
            inputs=[mode_input],
            outputs=[provider_input, model_input, base_url_input, api_key_input, connection_btn, connection_output, ai_selection_input, ai_confirm_input, ai_plan_output],
        )
        provider_input.change(
            update_provider_fields,
            inputs=[provider_input], outputs=[model_input, base_url_input, api_key_input],
        )
        connection_btn.click(
            test_provider_connection,
            inputs=[provider_input, model_input, base_url_input, api_key_input],
            outputs=[connection_output],
        )
        local_search_btn.click(
            search_local_index_for_ui,
            inputs=[local_query, retrieval_mode_input, search_context_state, semantic_ui_state],
            outputs=[local_search_output],
        )
        qa_btn.click(
            answer_evidence_only_for_ui,
            inputs=[qa_question, retrieval_mode_input, search_context_state, semantic_ui_state],
            outputs=[qa_output],
        )

        def update_retrieval_mode_visibility(mode):
            return gr.update(visible=mode == "Hybrid Local")

        def reset_semantic_for_ui_change(*_values):
            return reset_semantic_ui_state("模型 Profile 或本地模型目录已变化，请重新显式构建当前批次语义索引。")

        retrieval_mode_input.change(
            update_retrieval_mode_visibility,
            inputs=[retrieval_mode_input],
            outputs=[hybrid_settings],
        )
        model_profile_input.change(
            reset_semantic_for_ui_change,
            inputs=[model_profile_input, model_dir_input],
            outputs=[semantic_ui_state, semantic_status_output],
        )
        model_dir_input.change(
            reset_semantic_for_ui_change,
            inputs=[model_profile_input, model_dir_input],
            outputs=[semantic_ui_state, semantic_status_output],
        )
        build_semantic_btn.click(
            build_semantic_index_for_ui,
            inputs=[search_context_state, model_profile_input, model_dir_input],
            outputs=[semantic_ui_state, semantic_status_output],
        )

    return demo


if __name__ == "__main__":
    print("\n" + "="*55)
    print("  🧬 科研文献智能分析系统 — 网页版")
    print("="*55)
    print("  启动中，请稍候...")
    print("  启动后浏览器访问: http://localhost:7860")
    print("="*55 + "\n")

    demo = build_ui()
    demo.launch(
        server_name="127.0.0.1",
        server_port=7860,
        share=False,          # 改成 True 可生成公网链接（临时）
        inbrowser=True,       # 自动打开浏览器
        show_error=True,
    )
