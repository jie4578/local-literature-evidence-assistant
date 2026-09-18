import json
from pathlib import Path

import fitz
from docx import Document
from openpyxl import load_workbook

from batch_manager import resume_local_batch, run_local_batch
from exporters import export_excel, export_json, export_word
from local_summary import compare_papers
from report_modes import NO_WORD_MODE, REPORT_MODE_OPTIONS, report_layout_plan


def _paper(index: int) -> dict:
    return {
        "file_name": f"paper_{index}.pdf",
        "title_candidate": f"Study {index}",
        "author_candidate": f"Author {index}",
        "year_candidate": 2024,
        "doi": f"10.1234/study.{index}",
        "research_object": "adults",
        "sample_size": f"{index + 10} participants",
        "facts": [],
        "sections": {},
        "extractive_summary": {"evidence": [], "sentences": []},
    }


def _write_pdf(path: Path, text: str = "Abstract\nA local study.") -> Path:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


def test_report_modes_cover_requested_corpus_sizes():
    assert REPORT_MODE_OPTIONS == ("自动选择", "仅单篇报告", "简洁对比", "文献库汇总", NO_WORD_MODE)
    assert report_layout_plan("自动选择", 1)["layout"] == "single"
    assert report_layout_plan("自动选择", 2)["layout"] == "matrix"
    assert report_layout_plan("自动选择", 3)["layout"] == "matrix"
    for count in (4, 10, 30):
        plan = report_layout_plan("自动选择", count)
        assert plan["layout"] == "vertical"
        assert plan["max_word_papers"] == min(count, 10)
        assert plan["individual_reports"] is True
    assert report_layout_plan(NO_WORD_MODE, 30)["word_enabled"] is False


def test_thirty_papers_use_vertical_word_summary_and_reread_exports(tmp_path):
    papers = [_paper(index) for index in range(1, 31)]
    rows = compare_papers(papers)
    json_path = export_json(rows, tmp_path / "library.json")
    xlsx_path = export_excel(rows, tmp_path / "library.xlsx")
    docx_path = export_word(
        rows,
        tmp_path / "library.docx",
        papers=papers,
        report_layout_mode="自动选择",
        full_data_filename="library.xlsx / local_papers.json",
    )
    assert len(json.loads(json_path.read_text(encoding="utf-8"))) == 30
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook.active.max_row == 31
    document = Document(docx_path)
    assert all(len(table.columns) <= 4 for table in document.tables)
    assert any(len(table.rows) == 31 for table in document.tables)
    vertical_text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert all(f"paper_{index}.pdf" in vertical_text for index in range(1, 31))
    # 主报告只展开代表论文，完整库仍通过结构化文件保留。
    assert "完整 30 篇数据" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_local_batch_deduplicates_pause_resume_and_writes_manifest(tmp_path, monkeypatch):
    paths = [_write_pdf(tmp_path / f"重复测试_{index}.pdf", f"Page {index}") for index in range(1, 4)]
    calls = []

    def fake_extract(path):
        calls.append(Path(path).name)
        return {"file_name": Path(path).name, "pages": [], "facts": [], "sections": {}, "warnings": [], "errors": []}

    monkeypatch.setattr("batch_manager.extract_local_paper", fake_extract)
    first = run_local_batch([paths[0], paths[0], paths[1], paths[2]], output_root=tmp_path / "output", pause_after=1)
    assert first.manifest["status"] == "paused"
    assert len(first.papers) == 1
    manifest_path = first.task_dir / "batch_manifest.json"
    assert manifest_path.is_file()
    assert (first.task_dir / "status.json").is_file()
    resumed = resume_local_batch(first.task_dir, paths)
    assert len(resumed.papers) == 3
    assert resumed.manifest["status"] == "completed"
    assert len(resumed.manifest["papers"]) == 3
    assert len(calls) == 3  # 已完成的第一篇由缓存复用
    result_files = [first.task_dir / item["result_file"] for item in resumed.manifest["papers"]]
    assert all(path.is_file() for path in result_files)
    saved_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert saved_manifest["completed_count"] == 3


def test_local_batch_single_failure_keeps_other_results(tmp_path, monkeypatch):
    paths = [_write_pdf(tmp_path / "ok.pdf"), _write_pdf(tmp_path / "bad.pdf")]

    def fake_extract(path):
        if Path(path).name == "bad.pdf":
            return {"file_name": "bad.pdf", "pages": [], "facts": [], "sections": {}, "warnings": [], "errors": ["测试错误"]}
        return {"file_name": "ok.pdf", "pages": [], "facts": [], "sections": {}, "warnings": [], "errors": []}

    monkeypatch.setattr("batch_manager.extract_local_paper", fake_extract)
    result = run_local_batch(paths, output_root=tmp_path / "output")
    assert len(result.papers) == 2
    assert result.manifest["status"] == "completed_with_errors"
    assert result.manifest["completed_count"] == 1
    assert result.manifest["failed_count"] == 1
    assert all((result.task_dir / entry["result_file"]).is_file() for entry in result.manifest["papers"])


def test_changed_file_updates_cache_entry_instead_of_duplicating(tmp_path, monkeypatch):
    path = _write_pdf(tmp_path / "versioned.pdf", "v1")
    calls = []

    def fake_extract(current):
        calls.append(Path(current).read_bytes())
        return {"file_name": "versioned.pdf", "pages": [], "facts": [], "sections": {}, "warnings": [], "errors": []}

    monkeypatch.setattr("batch_manager.extract_local_paper", fake_extract)
    first = run_local_batch([path], output_root=tmp_path / "output")
    _write_pdf(path, "v2")
    second = resume_local_batch(first.task_dir, [path])
    assert len(second.manifest["papers"]) == 1
    assert len(calls) == 2


def test_ai_batch_requires_selection_and_confirmation_before_client(tmp_path, monkeypatch):
    import paper_claude

    path = _write_pdf(tmp_path / "candidate.pdf")
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test-only-placeholder")
    monkeypatch.setattr(paper_claude, "get_client", lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("尚未确认时不应创建 Provider")))
    result, output = paper_claude.process_papers([str(path)], "test-only-placeholder", mode="AI Provider", selected_ai_files=["candidate.pdf"], ai_confirmed=False)
    assert "确认" in result and output is None


def test_ai_request_estimate_is_local_only(tmp_path, monkeypatch):
    import paper_claude

    path = _write_pdf(tmp_path / "estimate.pdf", "Abstract\nA local estimate.")
    monkeypatch.setattr(paper_claude, "get_client", lambda *args, **kwargs: (_ for _ in ()).throw(AssertionError("估算阶段不应创建 Provider")))
    result = paper_claude.estimate_ai_plan([str(path)], ["estimate.pdf"], "DeepSeek", "deepseek-chat")
    assert "尚未创建 Provider" in result
    assert "预计" in result and "输出预算上限" in result
