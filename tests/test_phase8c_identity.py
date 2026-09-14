import json
from pathlib import Path
from zipfile import ZipFile

import fitz
import pytest
from docx import Document
from openpyxl import load_workbook

import storage
from batch_manager import MANIFEST_COMPATIBILITY_ERROR, resume_local_batch, run_local_batch
from exporters import export_excel, export_individual_reports, export_json, export_word
from local_summary import compare_papers
from report_modes import SELECTION_STRATEGY, select_word_papers


def _pdf(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()
    return path


def _fake_extract(path: str | Path) -> dict:
    source = Path(path)
    marker = source.read_text(encoding="utf-8", errors="ignore")
    if marker == "CORRUPT":
        raise ValueError("损坏 PDF")
    return {
        "file_name": source.name,
        "pages": [],
        "facts": [{"category": "sample_size", "verified": True}],
        "sections": {"abstract": {"text": marker}},
        "warnings": [],
        "errors": [],
        "marker": marker,
    }


def test_same_basename_different_content_gets_distinct_documents(tmp_path):
    first = tmp_path / "group_a" / "paper.pdf"
    second = tmp_path / "group_b" / "paper.pdf"
    first.parent.mkdir(parents=True, exist_ok=True)
    second.parent.mkdir(parents=True, exist_ok=True)
    first.write_text("UNIQUE-A", encoding="utf-8")
    second.write_text("UNIQUE-B", encoding="utf-8")
    result = run_local_batch([first, second], output_root=tmp_path / "output", extractor=_fake_extract)
    manifest = result.manifest
    assert len(manifest["inputs"]) == 2
    assert len(manifest["documents"]) == 2
    assert manifest["inputs"][0]["document_id"] != manifest["inputs"][1]["document_id"]
    assert {paper["marker"] for paper in result.papers} == {"UNIQUE-A", "UNIQUE-B"}


def test_same_content_different_filename_is_one_document_and_two_inputs(tmp_path):
    first = tmp_path / "first.pdf"
    second = tmp_path / "renamed.pdf"
    first.write_text("SAME-CONTENT", encoding="utf-8")
    second.write_text("SAME-CONTENT", encoding="utf-8")
    result = run_local_batch([first, second], output_root=tmp_path / "output", extractor=_fake_extract)
    assert len(result.manifest["inputs"]) == 2
    assert len(result.manifest["documents"]) == 1
    assert result.manifest["duplicate_count"] == 1
    assert result.manifest["inputs"][1]["status"] == "duplicate"
    assert result.manifest["inputs"][1]["duplicate_of"] == result.manifest["inputs"][0]["input_id"]
    assert len(result.papers) == 1


def test_same_filename_content_change_gets_new_document_identity(tmp_path):
    path = tmp_path / "versioned.pdf"
    path.write_text("VERSION-1", encoding="utf-8")
    calls = []

    def extract(current):
        calls.append(Path(current).read_text(encoding="utf-8"))
        return _fake_extract(current)

    first = run_local_batch([path], output_root=tmp_path / "output", extractor=extract)
    old_id = first.manifest["documents"][0]["document_id"]
    path.write_text("VERSION-2", encoding="utf-8")
    second = resume_local_batch(first.task_dir, [path], extractor=extract)
    assert len(second.manifest["documents"]) == 1
    assert second.manifest["documents"][0]["document_id"] != old_id
    assert second.papers[0]["marker"] == "VERSION-2"
    assert calls == ["VERSION-1", "VERSION-2"]


def test_corrupt_document_is_retained_but_not_reported(tmp_path):
    good = tmp_path / "good" / "paper.pdf"
    bad = tmp_path / "bad" / "paper.pdf"
    good.parent.mkdir(parents=True, exist_ok=True)
    bad.parent.mkdir(parents=True, exist_ok=True)
    good.write_text("GOOD", encoding="utf-8")
    bad.write_text("CORRUPT", encoding="utf-8")
    result = run_local_batch([good, bad], output_root=tmp_path / "output", extractor=_fake_extract)
    assert result.manifest["document_count"] == 2
    assert result.manifest["completed_count"] == 1
    assert result.manifest["failed_count"] == 1
    assert len(result.papers) == 2
    rows = compare_papers(result.papers)
    report_paths = export_individual_reports(rows, result.papers, result.task_dir / "individual_reports")
    assert len(report_paths) == 1
    assert all("paper_" not in path.name or path.is_file() for path in report_paths)


def test_pause_resume_keeps_input_mapping_and_avoids_reprocessing(tmp_path):
    paths = []
    for index in range(6):
        path = tmp_path / f"paper_{index}.pdf"
        path.write_text(f"DOC-{index}", encoding="utf-8")
        paths.append(path)
    calls = []

    def extract(current):
        calls.append(Path(current).name)
        return _fake_extract(current)

    first = run_local_batch(paths, output_root=tmp_path / "output", pause_after=3, extractor=extract)
    assert first.manifest["status"] == "paused"
    assert first.manifest["input_count"] == 6
    assert first.manifest["document_count"] == 6
    assert first.manifest["completed_count"] == 3
    second = resume_local_batch(first.task_dir, paths, extractor=extract)
    assert second.manifest["status"] == "completed"
    assert second.manifest["completed_count"] == 6
    assert len(second.manifest["inputs"]) == 6
    assert len(second.manifest["documents"]) == 6
    assert len(calls) == 6


def test_corrupt_document_cache_is_rebuilt_without_affecting_other_documents(tmp_path):
    first = tmp_path / "first.pdf"
    second = tmp_path / "second.pdf"
    first.write_text("CACHE-A", encoding="utf-8")
    second.write_text("CACHE-B", encoding="utf-8")
    calls = []

    def extract(current):
        calls.append(Path(current).name)
        return _fake_extract(current)

    initial = run_local_batch([first, second], output_root=tmp_path / "output", extractor=extract)
    first_cache = initial.task_dir / initial.manifest["documents"][0]["result_file"]
    first_cache.write_text("{broken json", encoding="utf-8")
    resumed = resume_local_batch(initial.task_dir, [first, second], extractor=extract)
    assert resumed.manifest["completed_count"] == 2
    assert calls == ["first.pdf", "second.pdf", "first.pdf"]
    assert json.loads(first_cache.read_text(encoding="utf-8"))["marker"] == "CACHE-A"


def test_atomic_write_preserves_old_json_and_cleans_temp_on_replace_failure(tmp_path, monkeypatch):
    target = tmp_path / "state.json"
    storage.atomic_write_json({"version": 1}, target)
    old = target.read_text(encoding="utf-8")

    def fail_replace(*args, **kwargs):
        raise OSError("simulated replace failure")

    monkeypatch.setattr(storage.os, "replace", fail_replace)
    with pytest.raises(OSError):
        storage.atomic_write_json({"version": 2}, target)
    assert target.read_text(encoding="utf-8") == old
    assert list(tmp_path.glob(".state.json.*.tmp")) == []


def test_old_manifest_is_rejected_without_guessing_migration(tmp_path):
    task = tmp_path / "old"
    task.mkdir()
    (task / "batch_manifest.json").write_text(json.dumps({"papers": []}), encoding="utf-8")
    with pytest.raises(ValueError, match="无法安全迁移"):
        resume_local_batch(task, [])
    assert MANIFEST_COMPATIBILITY_ERROR.startswith("批次 manifest")


def _selection_paper(index: int, complete: bool = True) -> dict:
    return {
        "file_name": f"paper_{index}.pdf",
        "document_id": f"doc_{index:064x}",
        "title_candidate": f"Study {index}",
        "author_candidate": f"Author {index}",
        "year_candidate": 2024,
        "doi": f"10.1234/study.{index}",
        "sections": {"abstract": {"text": "abstract"}, "methods": {"text": "methods"}},
        "facts": [{"verified": True}] * (index % 4),
        "text_quality_flags": [],
        "errors": [] if complete else ["失败"],
    }


def test_word_selection_is_deterministic_and_records_reason(tmp_path):
    papers = [_selection_paper(index) for index in range(30)]
    first = select_word_papers(papers)
    second = select_word_papers(list(reversed(papers)))
    first_ids = sorted(item["document_id"] for item in first if item["included_in_word"])
    second_ids = sorted(item["document_id"] for item in second if item["included_in_word"])
    assert len(first_ids) == 10
    assert first_ids == second_ids
    assert all(item["selection_strategy"] == SELECTION_STRATEGY for item in first)
    assert all(item["selection_reason"] for item in first)


def test_exports_keep_selection_and_document_identity_without_failed_report(tmp_path):
    papers = [_selection_paper(index) for index in range(3)] + [_selection_paper(99, complete=False)]
    for paper in papers:
        paper.update(select_word_papers([paper])[0])
    rows = compare_papers(papers)
    json_path = export_json(rows, tmp_path / "comparison.json")
    inputs = [
        {"input_id": "input_1", "display_name": "one.pdf", "source_alias": "one.pdf", "content_sha256": "abc", "document_id": papers[0]["document_id"], "status": "completed", "duplicate_of": None},
        {"input_id": "input_2", "display_name": "failed.pdf", "source_alias": "failed.pdf", "content_sha256": "def", "document_id": papers[-1]["document_id"], "status": "failed", "duplicate_of": None},
    ]
    xlsx_path = export_excel(rows, tmp_path / "comparison.xlsx", inputs=inputs, documents=[dict(paper, input_ids=[f"input_{index + 1}"], result_file=f"documents/{index}.json", status="completed" if not paper.get("errors") else "failed") for index, paper in enumerate(papers)])
    docx_path = export_word(rows, tmp_path / "comparison.docx", papers=papers, full_data_filename="comparison.xlsx / comparison.json")
    assert all("selection_reason" in row for row in json.loads(json_path.read_text(encoding="utf-8")))
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook.sheetnames == ["文献汇总", "输入映射", "技术明细"]
    assert workbook["文献汇总"].max_row == 5
    assert workbook["输入映射"].max_row == 3
    assert workbook["技术明细"].max_row == 5
    structured = load_workbook(xlsx_path, read_only=False)
    user_headers = [cell.value for cell in structured["文献汇总"][1]]
    assert all("_" not in str(value) for value in user_headers)
    assert structured["输入映射"]["F2"].value == "已完成"
    technical_headers = [cell.value for cell in structured["技术明细"][1]]
    assert "content_sha256" in technical_headers and "Sheet" not in structured.sheetnames
    document = Document(docx_path)
    assert all(len(table.columns) <= 4 for table in document.tables)
    reports = export_individual_reports(rows, papers, tmp_path / "individual_reports")
    assert len(reports) == 3
    assert all(paper["document_id"][:12] in path.name for paper, path in zip(papers[:3], reports))


def test_word_missing_fields_use_chinese_labels(tmp_path):
    paper = {
        "file_name": "missing.pdf",
        "title_candidate": "A study",
        "facts": [],
        "extractive_summary": {"evidence": []},
    }
    row = compare_papers([paper])[0]
    path = export_word([row], tmp_path / "missing.docx", papers=[paper])
    text = "\n".join(paragraph.text for paragraph in Document(path).paragraphs)
    text += "\n" + "\n".join(cell.text for table in Document(path).tables for row in table.rows for cell in row.cells)
    assert "第一作者" in text and "研究对象" in text and "样本量" in text
    assert "author_candidate" not in text and "research_object" not in text and "sample_size" not in text


def test_nfkc_equivalent_micro_sign_keeps_raw_evidence_verified():
    from scientific_facts import extract_scientific_facts

    pages = [{
        "source_file": "unicode.pdf",
        "page_number": 1,
        "raw_text": "The exposure was 0.6 µg/mL.",
        "display_text": "The exposure was 0.6 μg/mL.",
    }]
    facts = extract_scientific_facts(pages)
    concentration = [fact for fact in facts if fact["category"] == "concentration"]
    assert concentration and concentration[0]["verified"] is True
    assert "µ" in concentration[0]["evidence_quote_raw"]
    assert "μ" in concentration[0]["evidence_quote_display"]
