import json
import re
from pathlib import Path
from zipfile import ZipFile

import fitz
import pytest
from docx import Document
from openpyxl import load_workbook

from exporters import DISPLAY_LABELS, _conclusions, _limitations, _major_results, export_csv, export_excel, export_json, export_word
from bilingual import translate_paper_for_display
from local_extractor import _clean_page_text, extract_local_paper
from local_search import LocalSearchIndex
from local_summary import compare_papers, extractive_summary
from paper_pipeline import PageText
from section_parser import parse_sections
from scientific_facts import extract_scientific_facts, repair_visual_word_breaks


def make_local_pdf(path: Path) -> Path:
    doc = fitz.open()
    pages = [
        "A Controlled Study\nAlice Author\n2024\ndoi: 10.1234/example.doi\nAbstract\nWe evaluated stability.",
        "Research purpose: Evaluate short-term stability.\nMethods\nThe study included 120 independent samples.\nThe randomized, double-blind, controlled design stored samples at 4 degrees Celsius for eight weeks.",
        "Results\nThe result was statistically significant with p = 0.003.\nThe mean was 92.4 ± 2.1.",
        "Discussion\nThe limitation was a single laboratory and short duration.\nConclusion\nABX-17 showed greater stability.",
        "References\nA. Author. 2024.",
    ]
    for text in pages:
        page = doc.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    doc.save(path)
    doc.close()
    return path


def test_local_extraction_sections_metadata_and_facts(tmp_path):
    result = extract_local_paper(make_local_pdf(tmp_path / "中文论文.pdf"))
    assert len(result["pages"]) == 5
    assert result["title_candidate"] == "A Controlled Study"
    assert result["doi"] == "10.1234/example.doi"
    assert "abstract" in result["sections"] and "methods" in result["sections"]
    categories = {fact["category"] for fact in result["facts"]}
    assert {"sample_size", "temperature", "time", "p_value", "mean_sd", "randomization", "double_blind", "limitation", "conclusion"} <= categories
    assert all(fact["verified"] is True for fact in result["facts"])
    assert all(fact["pdf_page_start"] == fact["pdf_page_end"] for fact in result["facts"])


def test_section_parser_requires_explicit_heading():
    pages = [PageText("a.pdf", 1, "The results were discussed.\nNot a heading", 0, False)]
    assert parse_sections(pages) == {}


def test_section_parser_recognizes_introduction_variants_and_background_heading():
    pages = [PageText("a.pdf", 1, "1. Introduction\nintro\n2. Methods\nmethod", 0, False)]
    sections = parse_sections(pages)
    assert "introduction" in sections and "methods" in sections
    background = parse_sections([PageText("a.pdf", 1, "Background\ncontext", 0, False)])
    assert "introduction" in background


def test_section_parser_stops_at_author_and_publication_sections():
    pages = [PageText("a.pdf", 1, "Conclusion\nThe intervention improved outcomes.\nAuthors' contributions\nCI (and AON) developed the bloods protocol for the RCT.\nCompeting interests\nFunding details.", 0, False)]
    sections = parse_sections(pages)
    assert "CI (and AON) developed" not in sections["conclusion"]["text"]
    assert "Funding details" not in sections["conclusion"]["text"]


def test_all_excluded_publication_headings_are_boundaries():
    headings = [
        "Authors' contributions", "Author contributions", "Acknowledgements", "Funding",
        "Availability of data and materials", "Ethics approval", "Consent for publication",
        "Competing interests", "Publisher's Note",
    ]
    for heading in headings:
        sections = parse_sections([PageText("a.pdf", 1, f"Conclusion\nA conclusion.\n{heading}\nExcluded text.", 0, False)])
        assert "Excluded text" not in sections["conclusion"]["text"]


def test_local_search_fts_duplicate_update_phrase_and_clear(tmp_path):
    pages = [PageText("a.pdf", 1, "alpha beta phrase", 17, False), PageText("a.pdf", 2, "gamma", 5, False)]
    index = LocalSearchIndex(tmp_path / "index.sqlite")
    assert index.index_pages(pages) == 2
    assert index.index_pages(pages) == 0
    assert index.search('"alpha beta"')[0]["page_number"] == 1
    assert index.search("alpha gamma", limit=5) == []
    updated = [PageText("a.pdf", 1, "updated text", 12, False)]
    assert index.index_pages(updated) == 1
    assert index.search("updated")[0]["snippet"] == "updated text"
    with pytest.raises(ValueError):
        index.clear()
    index.clear(confirm=True)


def test_local_search_supports_statistical_p_value_variants(tmp_path):
    pages = [PageText("stats.pdf", 1, "The p = 0.003 and P<0.05 findings were significant.", 48, False),
             PageText("stats.pdf", 2, "The p-value was reported with p\u202f<\u202f0.001.", 45, False),
             PageText("stats.pdf", 3, "统计结果 p值显著。", 12, False)]
    index = LocalSearchIndex(tmp_path / "stats.sqlite")
    index.index_pages(pages)
    assert index.search("p value")
    assert index.search("p-value")
    assert index.search("p值")
    assert index.search("p = 0.003")
    assert index.search("p < 0.001")
    assert index.search("P<0.05")
    assert index.search("alpha") == []


def test_offline_summary_and_exports_are_structured(tmp_path):
    paper = extract_local_paper(make_local_pdf(tmp_path / "paper.pdf"))
    rows = compare_papers([paper])
    assert rows[0]["file_name"] == "paper.pdf"
    summary = extractive_summary(paper)
    assert summary["summary_type"] == "extractive"
    assert summary["description"].startswith("摘取式摘要")
    assert all(sentence in "\n".join(page["text"] for page in paper["pages"]) for sentence in summary["sentences"])
    json_path = export_json(rows, tmp_path / "中文.json")
    csv_path = export_csv(rows, tmp_path / "comparison.csv")
    xlsx_path = export_excel(rows, tmp_path / "comparison.xlsx")
    docx_path = export_word(rows, tmp_path / "comparison.docx")
    assert json.loads(json_path.read_text(encoding="utf-8"))[0]["file_name"] == "paper.pdf"
    assert all(path.exists() for path in (csv_path, xlsx_path, docx_path))


def test_extractive_summary_uses_sections_and_excludes_front_matter():
    paper = {
        "file_name": "paper.pdf",
        "pages": [{"page_number": 1, "text": "RESEARCH ARTICLE\nOpen Access\nSplit\nTitle\nAlice Author\nUniversity\nAbstract\nWe evaluated the treatment purpose in adults."},
                   {"page_number": 2, "text": "Methods\nThis randomized controlled trial included 120 participants.\nResults\nThe result was significant with p = 0.003.\nConclusion\nThe treatment improved outcomes."}],
        "sections": {
            "abstract": {"title": "Abstract", "text": "We evaluated the treatment purpose in adults.", "page_start": 1, "page_end": 1},
            "methods": {"title": "Methods", "text": "This randomized controlled trial included 120 participants.", "page_start": 2, "page_end": 2},
            "results": {"title": "Results", "text": "The result was significant with p = 0.003.", "page_start": 2, "page_end": 2},
            "conclusion": {"title": "Conclusion", "text": "The treatment improved outcomes.", "page_start": 2, "page_end": 2},
        },
    }
    summary = extractive_summary(paper)
    joined = " ".join(summary["sentences"])
    assert "RESEARCH ARTICLE" not in joined and "Open Access" not in joined
    assert "Alice Author" not in joined and "University" not in joined
    assert any("purpose" in sentence for sentence in summary["sentences"])
    assert any("randomized" in sentence for sentence in summary["sentences"])
    assert any("p = 0.003" in sentence for sentence in summary["sentences"])
    assert all(item["verified"] and item["pdf_pages"] for item in summary["evidence"])


def test_publication_front_matter_is_not_rule_evidence():
    pages = [{
        "source_file": "paper.pdf", "page_number": 1,
        "raw_text": "RESEARCH ARTICLE\nOpen Access\nA randomised controlled trial of diet improvement.\nFelice N. Jacka\nAbstract\nWe randomised 67 participants.",
        "text": "RESEARCH ARTICLE\nOpen Access\nA randomised controlled trial of diet improvement.\nFelice N. Jacka\nAbstract\nWe randomised 67 participants.",
        "display_text": "RESEARCH ARTICLE\nOpen Access\nA randomised controlled trial of diet improvement.\nFelice N. Jacka\nAbstract\nWe randomised 67 participants.",
    }]
    facts = extract_scientific_facts(pages)
    assert all("RESEARCH ARTICLE" not in fact["evidence_quote_display"] for fact in facts)
    assert all("Open Access" not in fact["evidence_quote_display"] for fact in facts)


def test_raw_and_display_evidence_are_separate_and_raw_is_exact():
    pages = [{"source_file": "a.pdf", "page_number": 1, "text": "admin-\nistration result", "raw_text": "admin-\nistration result", "display_text": "administration result"}]
    facts = extract_scientific_facts(pages)
    assert facts == [] or all(fact["evidence_quote_raw"] in pages[0]["raw_text"] for fact in facts)


def test_display_removes_controls_but_raw_is_preserved_and_flagged(tmp_path):
    display, flags = _clean_page_text("value\x01 with µ °C ± α β ≤ ≥", set())
    assert "\x01" not in display
    assert "raw_control_character_removed_for_display" in flags


def test_fact_with_section_heading_prefix_is_assigned_to_that_section():
    pages = [{
        "source_file": "heading.pdf", "page_number": 1,
        "raw_text": "Results\nThe observed response increased: R² = 0.92, α and β were ≤ 1.0 and ≥ 0.1, with p = 0.001.",
        "text": "Results\nThe observed response increased: R² = 0.92, α and β were ≤ 1.0 and ≥ 0.1, with p = 0.001.",
        "display_text": "Results\nThe observed response increased: R² = 0.92, α and β were ≤ 1.0 and ≥ 0.1, with p = 0.001.",
    }]
    facts = extract_scientific_facts(pages)
    sections = {"results": {"title": "Results", "text": "The observed response increased: R² = 0.92, α and β were ≤ 1.0 and ≥ 0.1, with p = 0.001.", "page_start": 1, "page_end": 1}}
    from local_extractor import _assign_fact_sections
    _assign_fact_sections(facts, sections)
    assert facts and all(fact["section"] == "results" for fact in facts)


def test_unicode_exports_round_trip(tmp_path):
    value = [{"evidence_quote_raw": "µ °C ± α β ≤ ≥ – − ²", "evidence_quote_display": "µ °C ± α β ≤ ≥ – − ²"}]
    json_path = export_json(value, tmp_path / "unicode.json")
    assert json.loads(json_path.read_text(encoding="utf-8"))[0]["evidence_quote_display"] == value[0]["evidence_quote_display"]
    csv_path = export_csv(value, tmp_path / "unicode.csv")
    assert "µ" in csv_path.read_text(encoding="utf-8-sig")
    xlsx_path = export_excel(value, tmp_path / "unicode.xlsx")
    assert load_workbook(xlsx_path, read_only=True).active.cell(2, 1).value == value[0]["evidence_quote_raw"]
    docx_path = export_word(value, tmp_path / "unicode.docx")
    assert "µ" in "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)


def test_safe_line_break_hyphen_cleaning_and_compound_preservation():
    display, flags = _clean_page_text("admin-\nistration\ndouble-blind\nsingle-dose\nLC-MS", set())
    assert "administration" in display
    assert "double-blind" in display and "single-dose" in display and "LC-MS" in display
    assert "unresolved_line_break_hyphen" not in flags


def test_unsafe_line_break_hyphen_is_not_guessed():
    pages = [{"source_file": "a.pdf", "page_number": 1,
              "raw_text": "The Modi-\nMedDiet result was observed.",
              "display_text": "The Modi-\nMedDiet result was observed.",
              "text": "The Modi-\nMedDiet result was observed."}]
    facts = extract_scientific_facts(pages)
    assert facts == [] or all("unresolved_line_break_hyphen" in fact.get("quality_flags", []) for fact in facts)


def test_local_mode_does_not_create_deepseek_client(monkeypatch, tmp_path):
    import paper_claude

    monkeypatch.setattr(paper_claude, "get_client", lambda *_: pytest.fail("Local Offline 不应创建 DeepSeek 客户端"))
    file_path = make_local_pdf(tmp_path / "offline.pdf")
    result, output = paper_claude.process_local_papers([str(file_path)])
    assert "Local Offline" in result
    assert output and Path(output).exists()


def test_research_object_is_conservative_and_keeps_page_evidence(tmp_path):
    path = tmp_path / "objects.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "A monkey pharmacokinetic (PK) studies report on MAB1.")
    doc.new_page().insert_text((72, 72), "Adults with major depression joined the trial.")
    doc.save(path)
    doc.close()
    first = extract_local_paper(path)
    assert first["research_object"] == "monkey pharmacokinetic (PK) studies"
    assert first["research_object_pdf_page"] == 1


def test_compare_repairs_cross_line_sample_and_deduplicates_methods():
    paper = {
        "file_name": "trial.pdf", "title_candidate": "Trial", "author_candidate": "Author",
        "year_candidate": 2024, "doi": "10.1234/trial", "research_object": "adults",
        "facts": [
            {"category": "sample_size", "evidence_quote_display": "We enrolled 67 participants (intervention, n = 33; control,", "pdf_page_start": 1, "pdf_page_end": 1, "source_type": "unknown"},
            {"category": "sample_size", "evidence_quote_display": "n = 34).", "pdf_page_start": 1, "pdf_page_end": 1, "source_type": "unknown"},
            {"category": "sample_size", "evidence_quote_display": "Table 1 Total (n = 67) intervention (n = 33) control (n = 34)", "pdf_page_start": 1, "pdf_page_end": 1, "source_type": "possible_table"},
            {"category": "randomization"}, {"category": "randomization"}, {"category": "control_group"},
        ],
        "sections": {"methods": {"text": "A randomized controlled trial."}},
    }
    row = compare_papers([paper])[0]
    assert "67" in row["sample_size"] and "n = 33" in row["sample_size"] and "n = 34" in row["sample_size"]
    assert "sample_size" not in row["missing_fields"]
    assert row["research_methods_keywords"].count("randomization") == 1


def test_summary_excludes_incomplete_fragments_and_keeps_verified_evidence():
    paper = {
        "pages": [{"page_number": 1, "raw_text": "Results\nComplete result was significant.\nThe dietary support group demonstrated significantly\nFigs 4–", "text": "Results\nComplete result was significant.\nThe dietary support group demonstrated significantly\nFigs 4–", "display_text": "Results\nComplete result was significant.\nThe dietary support group demonstrated significantly\nFigs 4–"}],
        "sections": {"results": {"title": "Results", "text": "Complete result was significant.\nThe dietary support group demonstrated significantly\nFigs 4–", "page_start": 1, "page_end": 1}},
    }
    summary = __import__("local_summary").extractive_summary(paper)
    assert summary["sentences"] == ["Complete result was significant."]
    assert summary["evidence"][0]["verified"] is True


class TranslationProvider:
    def __init__(self, response):
        self.response = response
        self.calls = []

    def generate(self, messages, **kwargs):
        self.calls.append(kwargs)
        from providers import ProviderResponse
        return ProviderResponse.from_legacy_string(self.response)


def test_translation_is_bounded_and_preserves_statistics():
    paper = {"title_candidate": "Trial", "sample_size": "67 participants, p = 0.003", "evidence": [{"evidence_quote_display": "67 participants, p = 0.003", "verified": True, "pdf_page_start": 2}]}
    provider = TranslationProvider('{"translations":{"title_candidate":"试验","sample_size":"67名参与者，p = 0.003","evidence_0":"67名参与者，p = 0.003"}}')
    outcome = translate_paper_for_display(provider, paper)
    assert outcome["status"] == "success"
    assert paper["sample_size_zh"] == "67名参与者，p = 0.003"
    assert paper["evidence"][0]["translation_status"] == "success"
    assert provider.calls[0]["max_output_tokens"] > 0


def test_translation_failure_keeps_english_fallback():
    paper = {"sample_size": "67 participants, p = 0.003", "evidence": [{"evidence_quote_display": "67 participants, p = 0.003", "verified": True}]}
    provider = TranslationProvider('{"translations":{"sample_size":"六十七名参与者"}}')
    outcome = translate_paper_for_display(provider, paper)
    assert outcome["status"] == "failed"
    assert "sample_size_zh" not in paper
    assert paper["evidence"][0]["translation_status"] == "failed"


def test_word_report_has_comparison_and_evidence_tables_without_internal_fields(tmp_path):
    paper = {"file_name": "trial.pdf", "title_candidate": "Trial", "author_candidate": "Author", "year_candidate": 2024, "doi": "10.1234/trial", "research_object": "adults", "sample_size": "67 participants", "facts": [{"category": "sample_size", "evidence_quote_raw": "67 participants", "evidence_quote_display": "67 participants", "pdf_page_start": 2, "pdf_page_end": 2, "verified": True, "source_type": "body", "quality_flags": []}], "extractive_summary": {"sentences": ["67 participants"]}}
    rows = compare_papers([paper])
    path = export_word(rows, tmp_path / "report.docx", papers=[paper])
    document = Document(path)
    xml_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert len(document.tables) >= 2
    assert "类型与页码" in xml_text and "PDF 第 2 页" in xml_text
    assert "research_object" not in xml_text and "None" not in xml_text and "[]" not in xml_text


def test_word_report_uses_compact_columns_and_merges_duplicate_evidence(tmp_path):
    shared = {"evidence_quote_raw": "The study included 67 participants.", "evidence_quote_display": "The study included 67 participants.", "pdf_page_start": 2, "pdf_page_end": 2, "verified": True, "source_type": "body", "quality_flags": []}
    first = {"file_name": "one.pdf", "title_candidate": "One", "facts": [dict(shared, category="sample_size"), dict(shared, category="major_result")], "extractive_summary": {"evidence": []}}
    second = {"file_name": "two.pdf", "title_candidate": "Two", "facts": [dict(shared, category="sample_size")], "extractive_summary": {"evidence": []}}
    rows = compare_papers([first, second])
    path = export_word(rows, tmp_path / "compact.docx", papers=[first, second])
    document = Document(path)
    assert len(document.tables[0].columns) == 3
    assert [len(table.columns) for table in document.tables[2::2]] == [4, 4]
    assert document.tables[2].rows.__len__() == 2  # header + one merged evidence row
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    assert b"w:wordWrap" in xml and b"w:eastAsia" in xml
    assert b"w:tblGrid" in xml and b"w:gridCol" in xml and b"w:cantSplit" in xml
    assert b'w:w="1866"' in xml and b'w:w="5910"' in xml


def test_fact_classification_and_display_word_repairs_are_conservative():
    pages = [{"source_file": "paper.pdf", "page_number": 3,
              "raw_text": "Model fitting criteria used a goodness of fit threshold. The biomarker i ncreased significantly. The control value remained unchanged. Results are preliminary and require replication in larger sample studies.",
              "text": "Model fitting criteria used a goodness of fit threshold. The biomarker i ncreased significantly. The control value remained unchanged. Results are preliminary and require replication in larger sample studies.",
              "display_text": "Model fitting criteria used a goodness of fit threshold. The biomarker i ncreased significantly. The control value remained unchanged. Results are preliminary and require replication in larger sample studies."}]
    facts = extract_scientific_facts(pages)
    by_category = {fact["category"]: fact for fact in facts}
    assert "research_method" in by_category
    assert "major_result" in by_category
    assert "limitation" in by_category
    major_results = [fact for fact in facts if fact["category"] == "major_result"]
    assert any("increased" in fact["evidence_quote_display"] for fact in major_results)
    assert any("i ncreased" in fact["evidence_quote_raw"] for fact in major_results)
    assert repair_visual_word_breaks("double-blind and single-dose") == "double-blind and single-dose"


def test_known_single_letter_display_breaks_are_repaired_without_changing_raw():
    value = "s ite i n t o r ecommendations 12-w eek s upport"
    assert repair_visual_word_breaks(value) == "site in to recommendations 12-week support"


def test_word_report_uses_chinese_labels_and_mode_specific_notice(tmp_path):
    evidence = [{"category": "time", "evidence_quote_raw": "The study lasted 8 weeks.", "evidence_quote_display": "The study lasted 8 weeks.", "pdf_page_start": 2, "pdf_page_end": 2, "verified": True, "source_type": "unknown", "quality_flags": []},
                {"category": "p", "evidence_quote_raw": "The result had p = 0.003.", "evidence_quote_display": "The result had p = 0.003.", "pdf_page_start": 3, "pdf_page_end": 3, "verified": True, "source_type": "unknown", "quality_flags": []},
                {"category": "conclusion", "evidence_quote_raw": "The treatment improved outcomes.", "evidence_quote_display": "The treatment improved outcomes.", "pdf_page_start": 4, "pdf_page_end": 4, "verified": True, "source_type": "unknown", "quality_flags": []}]
    paper = {"file_name": "trial.pdf", "title_candidate": "Trial", "facts": evidence, "extractive_summary": {"evidence": []}}
    rows = compare_papers([paper])
    local = Document(export_word(rows, tmp_path / "local.docx", papers=[paper], report_mode="Local Offline"))
    local_text = "\n".join([p.text for p in local.paragraphs] + [cell.text for table in local.tables for row in table.rows for cell in row.cells])
    assert "当前为 Local Offline 报告" in local_text
    assert "当前报告未启用语义翻译" not in local_text
    assert "时间条件" in local_text and "统计结果" in local_text and "研究结论" in local_text
    assert "source_type" not in local_text and "major_result" not in local_text

    translated = dict(paper)
    translated["facts"] = [dict(evidence[0], evidence_quote_zh="研究持续了8周。", translation_status="success")]
    translated_rows = compare_papers([translated])
    bilingual = Document(export_word(translated_rows, tmp_path / "bilingual.docx", papers=[translated], report_mode="中英对照"))
    bilingual_text = "\n".join([p.text for p in bilingual.paragraphs] + [cell.text for table in bilingual.tables for row in table.rows for cell in row.cells])
    assert "中文译文" in bilingual_text
    assert "当前报告包含可选机器翻译" in bilingual_text


def test_report_results_are_limited_to_complete_short_sentences():
    paper = {"extractive_summary": {"evidence": [
        {"section_key": "results", "text": "A complete result was significant.", "verified": True, "quality_flags": []},
        {"section_key": "results", "text": "The result was p < .001 (Fig.", "verified": True, "quality_flags": []},
        {"section_key": "results", "text": "The result was compared", "verified": True, "quality_flags": []},
        {"section_key": "results", "text": "A second complete result improved outcomes.", "verified": True, "quality_flags": []},
        {"section_key": "results", "text": "A third complete result remained stable.", "verified": True, "quality_flags": []},
        {"section_key": "results", "text": "A fourth complete result increased outcomes.", "verified": True, "quality_flags": []},
    ]}}
    values = _major_results(paper, {})
    assert len(values) == 3
    assert all(len(value) <= 300 and value.endswith(".") for value in values)
    assert all("Fig." not in value and not value.endswith("compared") for value in values)


def test_limitation_and_conclusion_selection_are_semantically_separate():
    paper = {"sections": {
        "discussion": {"text": "The oxidation level remained unchanged. The preliminary findings require replication in larger sample sizes."},
        "conclusion": {"text": "In summary, PQA monitoring supports product risk assessment."},
    }, "extractive_summary": {"evidence": []}}
    row = {"limitations_original": [], "conclusion_original": []}
    limitations = _limitations(paper, row)
    conclusion = _conclusions(paper, row)
    assert len(limitations) == 1 and "preliminary" in limitations[0]
    assert "oxidation" not in " ".join(limitations).lower()
    assert conclusion and "product risk assessment" in conclusion[0]
    assert set(limitations).isdisjoint(conclusion)


def test_word_final_spacer_is_one_point_with_zero_spacing(tmp_path):
    paper = {"file_name": "mock.pdf", "title_candidate": "Mock", "facts": [
        {"category": "sample_size", "evidence_quote_raw": "The study included 2 samples.", "evidence_quote_display": "The study included 2 samples.", "pdf_page_start": 1, "pdf_page_end": 1, "verified": True, "quality_flags": []},
    ], "extractive_summary": {"evidence": []}}
    path = export_word(compare_papers([paper]), tmp_path / "spacer.docx", papers=[paper])
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    paragraphs = re.findall(r"<w:p(?: [^>]*)?>.*?</w:p>", xml)
    assert paragraphs
    final = paragraphs[-1]
    assert 'w:line="20"' in final and 'w:lineRule="exact"' in final
    assert 'w:before="0"' in final and 'w:after="0"' in final
    assert 'w:sz w:val="2"' in final


def test_word_english_runs_use_en_us_and_disable_hyphenation(tmp_path):
    paper = {"file_name": "mock.pdf", "title_candidate": "Mock", "facts": [
        {"category": "p", "evidence_quote_raw": "The effects were significant, p = 0.003.", "evidence_quote_display": "The effects were significant, p = 0.003.", "pdf_page_start": 1, "pdf_page_end": 1, "verified": True, "quality_flags": []},
    ], "extractive_summary": {"evidence": []}}
    path = export_word(compare_papers([paper]), tmp_path / "language.docx", papers=[paper])
    document = Document(path)
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
    runs.extend(run for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs for run in paragraph.runs)
    for run in runs:
        if run.text and run.text.strip() and not re.search(r"[\u3400-\u9fff]", run.text):
            lang = run._element.get_or_add_rPr().find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}lang")
            assert lang is not None and lang.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val") == "en-US"
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert 'w:suppressAutoHyphens w:val="1"' in xml
    assert 'w:wordWrap w:val="0"' in xml


def test_display_type_mapping_keeps_method_distinct_from_model_fit():
    assert DISPLAY_LABELS["research_method"] == "研究方法"
    assert DISPLAY_LABELS["model_fit"] == "模型拟合"
    assert DISPLAY_LABELS["r_square"] == "模型拟合"
    assert DISPLAY_LABELS["temperature"] == "温度条件"


def test_r_square_is_method_evidence_not_major_result_and_word_label(tmp_path):
    pages = [{
        "source_file": "fit.pdf", "page_number": 1,
        "raw_text": "Methods\nR square = 0.92 and the model fit was acceptable.",
        "text": "Methods\nR square = 0.92 and the model fit was acceptable.",
        "display_text": "Methods\nR square = 0.92 and the model fit was acceptable.",
    }]
    facts = extract_scientific_facts(pages)
    fit_facts = [fact for fact in facts if "R square" in fact["evidence_quote_raw"]]
    assert fit_facts and all(fact["category"] in {"research_method", "model_fit"} for fact in fit_facts)
    assert not any(fact["category"] == "major_result" for fact in fit_facts)
    paper = {"file_name": "fit.pdf", "title_candidate": "Fit", "facts": fit_facts, "extractive_summary": {"evidence": []}}
    path = export_word(compare_papers([paper]), tmp_path / "fit.docx", papers=[paper])
    evidence_table = Document(path).tables[-1]
    evidence_text = "\n".join(cell.text for row in evidence_table.rows[1:] for cell in row.cells)
    assert "模型拟合" in evidence_text and "主要结果" not in evidence_text


def test_summary_excludes_author_contribution_section_content():
    raw = "Conclusion\nThe intervention improved outcomes.\nAuthors' contributions\nCI (and AON) developed the bloods protocol for the RCT."
    pages = [{"source_file": "smiles.pdf", "page_number": 1, "raw_text": raw, "text": raw, "display_text": raw}]
    sections = parse_sections(pages)
    paper = {"pages": pages, "sections": sections}
    summary = extractive_summary(paper)
    assert all("bloods protocol" not in item.get("text", "") for item in summary["evidence"])
    assert all("authors" not in item.get("text", "").casefold() for item in summary["evidence"])
    assert all(not item.get("text", "").casefold().startswith(("authors", "ci (and aon)")) for item in summary["evidence"])


def test_summary_and_conclusion_filter_publication_statements_and_questions():
    paper = {
        "pages": [{"page_number": 1, "raw_text": "Abstract\nA complete intervention improved outcomes.\nAll authors read and approved the manuscript.\nIf I improve my diet, will my mental health improve?\nThe intervention had significant effects.", "display_text": "Abstract\nA complete intervention improved outcomes.\nAll authors read and approved the manuscript.\nIf I improve my diet, will my mental health improve?\nThe intervention had significant effects."}],
        "sections": {
            "abstract": {"title": "Abstract", "text": "A complete intervention improved outcomes. All authors read and approved the manuscript.", "page_start": 1, "page_end": 1},
            "conclusion": {"title": "Conclusion", "text": "If I improve my diet, will my mental health improve? The intervention had significant effects.", "page_start": 1, "page_end": 1},
        },
    }
    summary = extractive_summary(paper)
    assert all("All authors read and approved" not in item["text"] for item in summary["evidence"])
    conclusion = _conclusions(paper, {})
    assert conclusion and not conclusion[0].endswith("?")
    assert "intervention" in conclusion[0].lower() or "effects" in conclusion[0].lower()
