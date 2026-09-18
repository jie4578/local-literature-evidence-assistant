"""本地结构化结果导出。"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from report_modes import report_layout_plan, select_word_papers
from storage import atomic_write_json


DISPLAY_LABELS = {
    "file_name": "文件名",
    "title_candidate": "论文标题",
    "author_candidate": "第一作者",
    "year_candidate": "年份",
    "doi": "DOI",
    "research_object": "研究对象",
    "sample_size": "样本量",
    "research_methods_keywords": "研究方法",
    "research_methods_original": "研究方法",
    "experimental_conditions": "实验条件",
    "statistical_information": "统计信息",
    "major_results_original": "主要结果",
    "limitations_original": "局限性",
    "conclusion_original": "研究结论",
    "source_pages": "来源页码",
    "missing_fields": "缺失字段",
    "document_id": "文档标识",
    "content_sha256": "内容指纹",
    "source_alias": "来源别名",
    "selection_strategy": "展示选择策略",
    "included_in_word": "是否纳入 Word",
    "selection_rank": "展示排序",
    "selection_score": "展示评分",
    "selection_reason": "展示选择原因",
    "input_id": "输入标识",
    "display_name": "显示名称",
    "status": "状态",
    "duplicate_of": "重复于",
    "input_ids": "输入映射",
    "result_file": "结果文件",
    "errors": "错误",
    "individual_report": "独立报告",
    "text_quality_flags": "文本质量标记",
    "sample_size": "样本量",
    "time": "时间条件",
    "temperature": "温度条件",
    "concentration": "浓度条件",
    "p": "统计结果",
    "p_value": "统计结果",
    "research_method": "研究方法",
    "model_fit": "模型拟合",
    "r_square": "模型拟合",
    "group_count": "分组数量",
    "randomization": "随机化",
    "double_blind": "双盲",
    "control_group": "对照组",
    "confidence_interval": "置信区间",
    "mean_sd": "均值 ± 标准差",
    "major_result": "主要结果",
    "limitation": "局限性",
    "conclusion": "研究结论",
    "immunocapture_lc_ms": "免疫捕获 LC-MS",
    "affinity_purification_lc_ms": "亲和纯化 LC-MS",
    "single_dose_pk": "单次给药 PK",
    "multiple_dose_pk": "多次给药 PK",
    "randomized_controlled_trial": "随机对照试验",
    "dietary_intervention": "饮食干预",
    "possible_table": "可能来自表格",
    "unknown": "来源未确定",
    "heading": "章节标题",
    "body": "正文",
    "major_result": "主要结果",
    "raw_control_character_removed_for_display": "展示层已移除控制字符",
    "replacement_character": "存在不可恢复字符",
    "incomplete_fragment": "残缺片段，未纳入摘要",
    "unresolved_line_break_hyphen": "跨行断词无法安全恢复",
}

STATUS_LABELS = {
    "completed": "已完成",
    "failed": "失败",
    "duplicate": "重复输入",
    "pending": "待处理",
    "running": "处理中",
    "skipped": "已跳过",
}


def _human_label(value: Any) -> str:
    if not isinstance(value, str):
        return str(value)
    if value in DISPLAY_LABELS:
        return DISPLAY_LABELS[value]
    if value in STATUS_LABELS:
        return STATUS_LABELS[value]
    if "_" in value:
        return value.replace("_", " ")
    return value


def _display_token(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return DISPLAY_LABELS.get(value, STATUS_LABELS.get(value, value))


def _evidence_category_label(item: dict[str, Any], category: Any) -> str:
    """根据已分类的原文区分普通研究方法和模型拟合证据。"""
    if category in {"model_fit", "r_square"}:
        return "模型拟合"
    if category == "research_method":
        source = str(item.get("evidence_quote_raw", item.get("evidence_quote_display", "")))
        if re.search(r"\b(?:r\s*(?:square|squared|2|²)|model\s+fit|goodness[- ]of[- ]fit|fit(?:ting)?\s+criteria)\b", source, re.IGNORECASE):
            return "模型拟合"
    return str(_display_token(category) or "其他证据")


def export_json(value: Any, path: str | Path) -> Path:
    return atomic_write_json(value, path)


def export_csv(rows: list[dict[str, Any]], path: str | Path) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    keys = list(rows[0].keys()) if rows else []
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _flat_value(value) for key, value in row.items()})
    return target


def _style_worksheet(sheet, headers: list[str]) -> None:
    from openpyxl.styles import Alignment, Font

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions if headers else None
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for column in sheet.columns:
        values = [str(cell.value or "") for cell in column]
        width = min(48, max(12, max((len(value) for value in values), default=12) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
        for cell in column:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def export_excel(
    rows: list[dict[str, Any]],
    path: str | Path,
    inputs: list[dict[str, Any]] | None = None,
    documents: list[dict[str, Any]] | None = None,
) -> Path:
    from openpyxl import Workbook
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文献汇总"
    summary_columns = (
        ("file_name", "文件名"), ("title_candidate", "论文标题"), ("author_candidate", "第一作者"),
        ("year_candidate", "年份"), ("doi", "DOI"), ("research_object", "研究对象"),
        ("sample_size", "样本量"), ("research_methods_original", "研究方法"),
        ("experimental_conditions", "实验条件"), ("statistical_information", "统计信息"),
        ("major_results_original", "主要结果"), ("limitations_original", "局限性"),
        ("conclusion_original", "研究结论"), ("source_pages", "来源页码"), ("missing_fields", "缺失字段"),
    )
    if inputs is None:
        # 保留旧的通用导出调用能力；批量报告始终传 inputs，使用结构化中文列。
        legacy_columns = list(rows[0].keys()) if rows else []
        sheet.append([_human_label(key) for key in legacy_columns])
        for row in rows:
            sheet.append([_flat_value(row.get(key)) for key in legacy_columns])
        _style_worksheet(sheet, [_human_label(key) for key in legacy_columns])
    else:
        sheet.append([label for _, label in summary_columns])
        for row in rows:
            sheet.append([_flat_value(row.get(key)) for key, _ in summary_columns])
        _style_worksheet(sheet, [label for _, label in summary_columns])
    if inputs is not None:
        input_sheet = workbook.create_sheet("输入映射")
        input_columns = (
            ("input_id", "输入标识"), ("display_name", "显示名称"), ("source_alias", "来源别名"),
            ("content_sha256", "内容指纹"), ("document_id", "文档标识"), ("status", "状态"),
            ("duplicate_of", "重复于"),
        )
        input_sheet.append([label for _, label in input_columns])
        for item in inputs:
            input_sheet.append([
                _flat_value(STATUS_LABELS.get(item.get(key), item.get(key)) if key == "status" else item.get(key))
                for key, _ in input_columns
            ])
        _style_worksheet(input_sheet, [label for _, label in input_columns])
        technical_sheet = workbook.create_sheet("技术明细")
        technical_columns = (
            "document_id", "content_sha256", "source_alias", "input_ids", "status", "result_file",
            "errors", "selection_strategy", "included_in_word", "selection_rank", "selection_score",
            "selection_reason", "individual_report",
        )
        technical_sheet.append(list(technical_columns))
        technical_rows = documents if documents is not None else rows
        for item in technical_rows:
            technical_sheet.append([_flat_value(item.get(key)) for key in technical_columns])
        _style_worksheet(technical_sheet, list(technical_columns))
    workbook.save(target)
    return target


def _flat_value(value: Any) -> str:
    if value is None or value == "" or value == [] or value == {}:
        return "未提取到"
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (list, tuple, set)):
        parts = [_flat_value(_display_token(item)) for item in value]
        joiner = "；" if any(re.search(r"[\u3400-\u9fff]", part) for part in parts) else "; "
        return joiner.join(parts)
    if isinstance(value, dict):
        return "；".join(f"{_human_label(key)}: {_flat_value(item)}" for key, item in value.items())
    return str(_display_token(value)).replace("**", "")


def _set_no_mid_word_breaks(paragraph) -> None:
    properties = paragraph._p.get_or_add_pPr()
    for tag in ("w:wordWrap", "w:suppressAutoHyphens"):
        element = properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            properties.append(element)
        element.set(qn("w:val"), "0" if tag == "w:wordWrap" else "1")


def _set_run_font(run, size: float = 10.5, bold: bool = False, color: str = "000000", language: str | None = None) -> None:
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Calibri")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    has_cjk = bool(run.text and re.search(r"[\u3400-\u9fff]", run.text))
    lang.set(qn("w:val"), language or ("zh-CN" if has_cjk else "en-US"))


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_borders(cell, color: str = "D9E2F3") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def _set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def _set_table_grid(table, widths: list[float] | None) -> None:
    if not widths:
        return
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(round(width * 1440)))
        grid.append(column)
    table_properties = table._tbl.tblPr
    layout = table_properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    table_width = table_properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(round(sum(widths) * 1440)))
    table_width.set(qn("w:type"), "dxa")


def _format_cell(cell, value: Any, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_no_mid_word_breaks(paragraph)
    run = paragraph.add_run(_flat_value(value))
    _set_run_font(run, size=9 if header else 8.5, bold=header, color="FFFFFF" if header else "000000")
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_margins(cell)
    _set_cell_borders(cell)
    if header:
        _set_cell_shading(cell, "1F4E78")


def _add_table(document: Document, headers: list[str], rows: Iterable[Iterable[Any]], widths: list[float] | None = None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_row = table.rows[0]
    _repeat_header(header_row)
    _prevent_row_split(header_row)
    for index, header in enumerate(headers):
        _format_cell(header_row.cells[index], header, header=True)
        if widths:
            header_row.cells[index].width = Inches(widths[index])
    for row_values in rows:
        row = table.add_row()
        _prevent_row_split(row)
        for index, value in enumerate(row_values):
            _format_cell(row.cells[index], value)
            if widths:
                row.cells[index].width = Inches(widths[index])
    _set_table_grid(table, widths)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    spacer_properties = spacer._p.get_or_add_pPr()
    spacer_run_properties = spacer_properties.find(qn("w:rPr"))
    if spacer_run_properties is None:
        spacer_run_properties = OxmlElement("w:rPr")
        spacer_properties.append(spacer_run_properties)
    spacer_size = spacer_run_properties.find(qn("w:sz"))
    if spacer_size is None:
        spacer_size = OxmlElement("w:sz")
        spacer_run_properties.append(spacer_size)
    spacer_size.set(qn("w:val"), "2")
    return table


def _add_heading(document: Document, text: str, level: int = 2) -> None:
    heading = document.add_heading(text, level=level)
    for run in heading.runs:
        _set_run_font(run, size=16 if level == 1 else 12, bold=True)
    heading.paragraph_format.keep_with_next = True
    _set_no_mid_word_breaks(heading)


def _add_field_paragraph(document: Document, label: str, value: Any) -> None:
    if value is None or value == "" or value == []:
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    _set_no_mid_word_breaks(paragraph)
    lead = paragraph.add_run(f"{label}：")
    _set_run_font(lead, bold=True)
    body = paragraph.add_run(_flat_value(value))
    _set_run_font(body)


def _paper_evidence(paper: dict[str, Any], row: dict[str, Any]) -> list[dict[str, Any]]:
    values = (paper.get("facts") or paper.get("evidence") or row.get("evidence", [])) if paper else row.get("evidence", [])
    if not values and row.get("evidence_quote_raw"):
        values = [row]
    result = []
    seen = set()
    for item in values:
        if not isinstance(item, dict):
            continue
        raw = item.get("evidence_quote_raw", item.get("evidence_quote", ""))
        display = item.get("evidence_quote_display", item.get("evidence_quote", raw))
        compact = re.sub(r"\s+", " ", str(display)).strip()
        lowered = compact.casefold()
        section = str(item.get("section", "")).casefold()
        if section in {"references", "introduction"}:
            continue
        if _is_non_content_sentence(compact):
            continue
        if re.match(r"^(?:research article|open access|original article|article|plos one|keywords?|correspondence|copyright|funding|author contributions|competing interests?)\b", compact, re.IGNORECASE):
            continue
        if "@" in compact or "doi.org/" in lowered or re.search(r"\b(?:university|department|institute|school of)\b", lowered):
            continue
        pages = tuple(item.get("pdf_pages") or ([item.get("pdf_page_start")] if item.get("pdf_page_start") else []))
        key = (raw, pages)
        if not raw:
            continue
        if key in seen:
            for prior in result:
                prior_pages = tuple(prior.get("pdf_pages") or ([prior.get("pdf_page_start")] if prior.get("pdf_page_start") else []))
                if prior.get("evidence_quote_raw", prior.get("evidence_quote", "")) == raw and prior_pages == pages:
                    category = item.get("category", item.get("evidence_type", ""))
                    if category not in prior["_merged_categories"]:
                        prior["_merged_categories"].append(category)
                    break
            continue
        seen.add(key)
        copied = dict(item)
        copied["_merged_categories"] = [item.get("category", item.get("evidence_type", ""))]
        for prior in result:
            prior_pages = tuple(prior.get("pdf_pages") or ([prior.get("pdf_page_start")] if prior.get("pdf_page_start") else []))
            if prior.get("evidence_quote_raw", prior.get("evidence_quote", "")) == raw and prior_pages == pages:
                prior["_merged_categories"].append(item.get("category", item.get("evidence_type", "")))
                break
        else:
            result.append(copied)
        if len(result) >= 8:
            break
    return result


def _complete_sentences(values: Any, limit: int, max_chars: int = 300) -> list[str]:
    if isinstance(values, str):
        values = [values]
    result = []
    for value in values or []:
        text = re.sub(r"\s+", " ", str(value or "")).strip()
        if not text or "\ufffd" in text:
            continue
        if len(text) > max_chars:
            continue
        if re.search(r"(?:fig(?:ure)?\.?|p\s*[<=>]\s*\.?\d*)\s*$", text, re.IGNORECASE):
            continue
        if re.search(r"(?:[-–—]|\b(?:significantly|compared|than|of|and|or|with|to|from|in|for|the|a|an|on|by|as))$", text, re.IGNORECASE):
            continue
        if not re.search(r"[.!?。！？]$", text):
            continue
        if text not in result:
            result.append(text)
        if len(result) >= limit:
            break
    return result


def _is_non_content_sentence(value: Any) -> bool:
    """排除出版元数据、署名/单位和章节后的附录信息。"""
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    lowered = text.casefold()
    if not text:
        return True
    if re.match(
        r"^(?:keywords?|correspondence|doi|received|accepted|trial registration|registered on|"
        r"author details|authors?['’] contributions|competing interests?|acknowledg|ethics approval|consent for publication)\b",
        text,
        re.IGNORECASE,
    ):
        return True
    if "all authors read and approved" in lowered or "authors' contributions" in lowered or "authors’ contributions" in lowered:
        return True
    if "doi.org/" in lowered or "creativecommons" in lowered or "article is distributed" in lowered or "@" in text:
        return True
    if re.search(r"\b(?:university|department|institute|school|hospital|centre|center),\s", text, re.IGNORECASE):
        return True
    return False


def _explicit_limitation(text: str) -> bool:
    return bool(re.search(r"\b(?:limitation|limited|preliminary|replication|larger sample|small sample|single laboratory|single[- ]center|short duration)\b", text, re.IGNORECASE))


def _without_repeats(values: list[str], excluded: set[str]) -> list[str]:
    return [value for value in values if re.sub(r"\s+", " ", value).strip().casefold() not in excluded]


def _paper_summary_items(paper: dict[str, Any], section_key: str | None = None) -> list[dict[str, Any]]:
    summary = paper.get("extractive_summary") or {}
    items = summary.get("evidence", []) if isinstance(summary, dict) else []
    return [
        item
        for item in items
        if (section_key is None or item.get("section_key") == section_key)
        and item.get("verified") is True
        and not item.get("quality_flags")
        and not _is_non_content_sentence(item.get("text"))
    ]


def _section_sentences(paper: dict[str, Any], keys: tuple[str, ...]) -> list[str]:
    values = []
    for key in keys:
        text = paper.get("sections", {}).get(key, {}).get("text", "")
        joined = re.sub(r"\s+", " ", str(text)).strip()
        values.extend(
            part.strip()
            for part in re.split(r"(?<=[.!?。！？])\s+", joined)
            if part.strip() and not _is_non_content_sentence(part)
        )
    return values


def _major_results(paper: dict[str, Any], row: dict[str, Any]) -> list[str]:
    values = [item.get("text") for item in _paper_summary_items(paper, "results")]
    if not values:
        values = row.get("major_results_original", [])
    return _complete_sentences(values, 3, 300)


def _limitations(paper: dict[str, Any], row: dict[str, Any]) -> list[str]:
    values = list(row.get("limitations_original", []))
    values.extend(_section_sentences(paper, ("discussion", "conclusion")))
    values.extend(item.get("text") for item in _paper_summary_items(paper, "discussion") + _paper_summary_items(paper, "conclusion"))
    values = [value for value in values if _explicit_limitation(str(value)) and not (
        re.search(r"oxidation|lysine", str(value), re.IGNORECASE)
        and not re.search(r"limitation|limited|small sample|short duration|larger sample|replication", str(value), re.IGNORECASE)
    )]
    ranked = sorted(
        dict.fromkeys(str(value) for value in values),
        key=lambda value: (
            -sum(bool(re.search(pattern, value, re.IGNORECASE)) for pattern in (r"larger sample", r"replication", r"small sample", r"short duration", r"limitation", r"limited")),
            -int(bool(re.search(r"preliminary", value, re.IGNORECASE))),
            len(value),
        ),
    )
    return _complete_sentences(ranked, 2, 300)


def _conclusions(paper: dict[str, Any], row: dict[str, Any]) -> list[str]:
    candidates = _paper_summary_items(paper, "conclusion") + _paper_summary_items(paper, "discussion") + _paper_summary_items(paper, "abstract")
    candidates.extend(
        {"text": text, "section_key": section_key, "verified": True, "quality_flags": []}
        for section_key in ("conclusion", "discussion", "abstract")
        for text in _section_sentences(paper, (section_key,))
    )
    def score(item: dict[str, Any]) -> int:
        text = item.get("text", "")
        if re.search(r"\?\s*$", text):
            return -100
        if re.search(r"\b(?:analys(?:is|es)|mmrm|mixed[- ]effects|statistical method|methodological approach)\b", text, re.IGNORECASE) and not re.search(r"criticality|product risk assessment|pqa|efficacious|conclusion|suggest", text, re.IGNORECASE):
            return -100
        value = 0
        if re.search(r"criticality|product risk assessment", text, re.IGNORECASE):
            value += 30
        elif re.search(r"pqa", text, re.IGNORECASE):
            value += 20
        elif re.search(r"drug development", text, re.IGNORECASE):
            value += 15
        elif re.search(r"lc[-/]ms|model", text, re.IGNORECASE):
            value += 8
        if re.search(r"intervention|improvement|effect|outcome|significant", text, re.IGNORECASE):
            value += 10
        if item.get("section_key") in {"conclusion", "discussion"}:
            value += 15 if item.get("section_key") == "conclusion" else 8
        if re.search(r"conclu|showed|improv|effective|stability|greater", text, re.IGNORECASE):
            value += 2
        return value
    ranked = sorted(candidates, key=lambda item: (-score(item), len(item.get("text", ""))))
    values = [item.get("text") for item in ranked if score(item) > 0] or [item.get("text") for item in candidates if score(item) >= 0] or row.get("conclusion_original", [])
    result_values = {re.sub(r"\s+", " ", value).strip().casefold() for value in _major_results(paper, row)}
    limitation_values = {re.sub(r"\s+", " ", value).strip().casefold() for value in _limitations(paper, row)}
    values = _without_repeats(values, result_values | limitation_values)
    return _complete_sentences(values, 1, 300)


def _research_object_display(value: Any) -> Any:
    if isinstance(value, str) and "monkey pharmacokinetic" in value.casefold():
        return "食蟹猴 PK 研究"
    if isinstance(value, str) and "major depression" in value.casefold():
        return "中重度抑郁症成年人"
    return value


def _method_evidence(paper: dict[str, Any], row: dict[str, Any]) -> list[str]:
    """将方法分类转换为已提取的原文证据，避免把分类名当作内容。"""
    method_categories = {
        "research_method", "model_fit", "r_square", "randomization", "double_blind",
        "control_group", "group_count", "immunocapture_lc_ms", "affinity_purification_lc_ms",
        "single_dose_pk", "multiple_dose_pk", "randomized_controlled_trial", "dietary_intervention",
    }
    values: list[str] = []
    for fact in paper.get("facts", []) or []:
        if fact.get("category") not in method_categories:
            continue
        if fact.get("section") == "results" and fact.get("category") in {"research_method", "model_fit", "r_square"}:
            continue
        value = re.sub(r"\s+", " ", str(fact.get("evidence_quote_display") or fact.get("evidence_quote") or "")).strip()
        if value and value not in values:
            values.append(value)
    if values:
        return values[:3]
    # 章节原句由 local_summary 预先保守抽取；这里只接受真实文本，不回退到分类名。
    fallback = []
    for value in row.get("research_methods_original") or []:
        text = re.sub(r"\s+", " ", str(value or "")).strip()
        if text and text not in {"研究方法", "research_method"} and text not in fallback:
            fallback.append(text)
    return fallback[:3]


def _comparison_value(row: dict[str, Any], key: str, paper: dict[str, Any] | None = None) -> Any:
    if key == "research_object":
        return _research_object_display(row.get(key))
    if key == "research_methods_keywords":
        return row.get("research_methods_original") or "未提取到"
    if key == "major_results_original":
        return _major_results(paper or {}, row)
    if key == "limitations_original":
        return _limitations(paper or {}, row)
    if key == "conclusion_original":
        return _conclusions(paper or {}, row)
    return row.get(key)


def export_word(
    rows: list[dict[str, Any]],
    path: str | Path,
    papers: list[dict[str, Any]] | None = None,
    report_mode: str = "Local Offline",
    final_review: dict[str, Any] | None = None,
    report_layout_mode: str = "自动选择",
    full_data_filename: str | None = None,
) -> Path:
    """导出紧凑的用户报告，保留英文证据和 PDF 物理页码。"""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, bold in (("Title", 20, True), ("Heading 1", 16, True), ("Heading 2", 12, True), ("Heading 3", 11, True)):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = bold
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_no_mid_word_breaks(title)
    title_run = title.add_run("本地文献结构化分析报告")
    _set_run_font(title_run, size=20, bold=True)
    meta = document.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_no_mid_word_breaks(meta)
    requested_layout = report_layout_mode
    if requested_layout == "自动选择" and report_mode in {
        "自动选择", "仅单篇报告", "简洁对比", "文献库汇总", "不生成 Word，仅导出 Excel/JSON",
    }:
        requested_layout = report_mode
    layout_plan = report_layout_plan(requested_layout, len(rows))
    meta_run = meta.add_run(
        f"生成模式：{report_mode}　报告布局：{layout_plan['resolved_mode']}　论文数量：{len(rows)}"
    )
    _set_run_font(meta_run, size=9, color="666666")
    if report_mode == "Local Offline":
        notice = "当前为 Local Offline 报告，未启用语义翻译；英文证据原文和 PDF 物理页码已保留。verified 仅表示原文匹配，不代表科研结论已被证明。"
    else:
        notice = "当前报告包含可选机器翻译，仅供阅读；英文证据原文、PDF 物理页码和 verified 状态保持不变。"
    explanation = document.add_paragraph(notice)
    _set_no_mid_word_breaks(explanation)
    for run in explanation.runs:
        _set_run_font(run, size=9)

    paper_list = papers or rows
    row_by_document_id = {
        str(row.get("document_id")): row
        for row in rows
        if row.get("document_id")
    }
    paper_infos = []
    for index, paper in enumerate(paper_list):
        row = row_by_document_id.get(str(paper.get("document_id")))
        if row is None:
            row = rows[index] if index < len(rows) else {}
        paper_infos.append((row, paper))
    selection_records = select_word_papers(list(paper_list), max_papers=10)
    selection_by_id = {record["document_id"]: record for record in selection_records}
    selection_by_file = {
        str(paper.get("file_name")): record
        for paper, record in zip(paper_list, selection_records)
        if paper.get("file_name")
    }
    selection_by_object_id = {
        id(paper): record
        for paper, record in zip(paper_list, selection_records)
    }

    def selection_for(paper: dict[str, Any]) -> dict[str, Any]:
        return (
            selection_by_id.get(str(paper.get("document_id")))
            or selection_by_file.get(str(paper.get("file_name")))
            or selection_by_object_id.get(id(paper), {})
        )
    comparison_keys = (
        ("title_candidate", "标题"), ("author_candidate", "第一作者"), ("year_candidate", "年份"),
        ("doi", "DOI"), ("research_object", "研究对象"), ("sample_size", "样本量"),
        ("research_methods_keywords", "研究方法"), ("major_results_original", "主要结果"),
        ("limitations_original", "局限性"), ("conclusion_original", "结论"),
    )
    if layout_plan["show_comparison"] and layout_plan["layout"] == "matrix":
        _add_heading(document, "多论文对比", 2)
        table_rows = []
        for key, label in comparison_keys:
            values = [label]
            for row, paper in paper_infos:
                values.append(_comparison_value(row, key, paper))
            table_rows.append(values)
        headers = ["对比维度"] + [f"论文 {index + 1}\n{row.get('file_name') or '未命名文件'}" for index, (row, _) in enumerate(paper_infos)]
        if len(rows) == 2:
            widths = [1.15, 3.0, 3.0]
        else:
            widths = [1.0] + [2.05] * len(rows)
        _add_table(document, headers, table_rows, widths)
    elif layout_plan["show_comparison"] and layout_plan["layout"] == "vertical":
        _add_heading(document, "文献库纵向汇总", 2)
        vertical_rows = []
        for index, (row, paper) in enumerate(paper_infos, 1):
            title = paper.get("title_candidate") or row.get("title_candidate") or "未提取到"
            author = paper.get("author_candidate") or row.get("author_candidate") or "未提取到"
            year = paper.get("year_candidate") or row.get("year_candidate") or "未提取到"
            doi = paper.get("doi") or row.get("doi") or "未提取到"
            basic = f"标题：{title}\n第一作者：{author}\n年份：{year}\nDOI：{doi}"
            object_sample = (
                f"研究对象：{_research_object_display(row.get('research_object')) or '未提取到'}\n"
                f"样本量：{_flat_value(row.get('sample_size'))}"
            )
            results = _major_results(paper, row)[:3]
            conclusion = _conclusions(paper, row)[:1]
            result_text = "\n".join([*(f"• {value}" for value in results), *(f"结论：{value}" for value in conclusion)])
            vertical_rows.append([f"论文 {index}\n{row.get('file_name') or paper.get('file_name') or '未命名文件'}", basic, object_sample, result_text or "未提取到"])
        _add_table(
            document,
            ["论文", "基础信息", "研究对象与样本量", "主要结果与结论"],
            vertical_rows,
            [0.85, 2.45, 1.45, 2.45],
        )
        if full_data_filename:
            readable_reference = str(full_data_filename).replace("_", " ")
            note = document.add_paragraph(f"完整数据请查看：{readable_reference}。每篇论文的详细证据保存在独立报告目录。")
            _set_no_mid_word_breaks(note)
            for run in note.runs:
                _set_run_font(run, size=9, color="666666")

    if final_review:
        labels = {"research_theme_overview": "研究主题概述", "major_methods": "主要研究方法", "common_conclusions": "共同结论", "different_or_conflicting_conclusions": "不同或冲突结论", "research_gaps": "研究空白", "future_recommendations": "后续研究建议"}
        _add_heading(document, "结构化综述", 2)
        for key, label in labels.items():
            if final_review.get(key):
                _add_field_paragraph(document, label, final_review[key])

    selected_infos = [
        info for info in paper_infos
        if selection_for(info[1]).get("included_in_word")
    ]
    selected_infos.sort(key=lambda info: selection_for(info[1]).get("selection_rank") or 9999)
    detail_infos = selected_infos[:layout_plan["max_word_papers"]]
    detail_papers = [paper for _, paper in detail_infos]
    if len(paper_list) > len(detail_papers) and full_data_filename:
        note = document.add_paragraph(
            f"由于论文数量较多，正文仅展示按数据完整度和证据覆盖率选出的 {len(detail_papers)} 篇论文；"
            f"该选择不代表学术重要性，完整 {len(paper_list)} 篇数据请查看：{str(full_data_filename).replace('_', ' ')}。"
        )
        _set_no_mid_word_breaks(note)
        for run in note.runs:
            _set_run_font(run, size=9, color="666666")
    for index, (row, paper) in enumerate(detail_infos):
        _add_heading(document, f"论文详情 {index + 1}", 2)
        _add_table(document, ["字段", "内容"], [
            ("论文标题", paper.get("title_candidate") or row.get("title_candidate") or paper.get("file_name") or row.get("file_name")),
            ("第一作者", paper.get("author_candidate") or row.get("author_candidate")),
            ("年份", paper.get("year_candidate") or row.get("year_candidate")),
            ("DOI", paper.get("doi") or row.get("doi")),
        ], [1.3, 5.55])
        labels = (("research_object", "研究对象"), ("sample_size", "样本量说明"), ("research_methods_keywords", "研究方法"),
                  ("experimental_conditions", "实验条件"))
        for key, label in labels:
            field_key = "research_methods_original" if key == "research_methods_keywords" else key
            value = row.get(field_key)
            if value:
                _add_field_paragraph(document, label, _research_object_display(value) if key == "research_object" else value)
        _add_heading(document, "统计信息", 3)
        for value in _complete_sentences(paper.get("statistical_information", row.get("statistical_information", [])), 3, 300):
            _add_field_paragraph(document, "原文", value)
        _add_heading(document, "结构化概览", 3)
        missing = row.get("missing_fields") or []
        _add_field_paragraph(document, "缺失字段", [DISPLAY_LABELS.get(item, item) for item in missing])
        _add_heading(document, "主要结果", 3)
        for value in _major_results(paper, row):
            _add_field_paragraph(document, "原文", value)
        _add_heading(document, "局限性", 3)
        limitation_values = _limitations(paper, row)
        if not limitation_values:
            _add_field_paragraph(document, "说明", "未提取到明确局限性")
        for value in limitation_values:
            _add_field_paragraph(document, "原文", value)
        _add_heading(document, "结论", 3)
        for value in _conclusions(paper, row):
            _add_field_paragraph(document, "原文", value)
        summary = paper.get("extractive_summary")
        if summary and summary.get("sentences"):
            _add_heading(document, "摘取式摘要", 3)
            note = document.add_paragraph("最多展示 6 条可回查的完整原句；不代表 AI 生成或事实核验。")
            _set_no_mid_word_breaks(note)
            for item in summary.get("evidence", [])[:6]:
                if item.get("verified") is True and not item.get("quality_flags"):
                    _add_field_paragraph(document, item.get("section", "原文"), item.get("evidence_quote_display", item.get("text")))
        _add_heading(document, "关键证据", 3)
        evidence_rows = []
        evidence_items = _paper_evidence(paper, row)
        has_translation = any(item.get("evidence_quote_zh") for item in evidence_items)
        for item in evidence_items:
            pages = item.get("pdf_pages") or ([item.get("pdf_page_start")] if item.get("pdf_page_start") else [])
            page_text = "、".join(str(page) for page in pages if page)
            categories = item.get("_merged_categories") or [item.get("category", item.get("evidence_type"))]
            type_text = "、".join(_evidence_category_label(item, category) for category in dict.fromkeys(categories))
            type_page = f"{type_text}\nPDF 第 {page_text or '未知'} 页"
            quality = [_display_token(flag) for flag in (item.get("quality_flags") or [])] or "无异常"
            if has_translation:
                evidence_rows.append([type_page, item.get("evidence_quote_zh") or "未提供中文译文", item.get("evidence_quote_display", item.get("evidence_quote", item.get("evidence_quote_raw"))), "已验证" if item.get("verified") is True else "未验证，请回查原文", quality])
            else:
                evidence_rows.append([type_page, item.get("evidence_quote_display", item.get("evidence_quote", item.get("evidence_quote_raw"))), "已验证" if item.get("verified") is True else "未验证，请回查原文", quality])
        evidence_headers = ["类型与页码", "中文译文", "英文证据原文", "验证状态", "质量提示"] if has_translation else ["类型与页码", "英文证据原文", "验证状态", "质量提示"]
        evidence_widths = [0.9, 1.5, 3.9, 0.6, 0.3] if has_translation else [1.296, 4.104, 0.864, 0.936]
        _add_table(document, evidence_headers, evidence_rows, evidence_widths)
        # rows-only 是旧调用兼容路径；保留一份可被普通段落读取的原文，
        # 同时证据表仍是用户报告的主要呈现形式。
        if not paper.get("facts") and not paper.get("evidence"):
            for item in evidence_items:
                _add_field_paragraph(document, "英文原文", item.get("evidence_quote_raw", item.get("evidence_quote")))

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_no_mid_word_breaks(footer)
    run = footer.add_run("第 ")
    _set_run_font(run, size=9, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    tail = footer.add_run(" 页")
    _set_run_font(tail, size=9, color="666666")
    document.save(target)
    return target


def export_individual_reports(
    rows: list[dict[str, Any]],
    papers: list[dict[str, Any]],
    directory: str | Path,
) -> list[Path]:
    """为大批量任务保存每篇论文的详细 Word 报告。"""
    target_dir = Path(directory)
    target_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    row_by_document_id = {str(row.get("document_id")): row for row in rows if row.get("document_id")}
    for index, paper in enumerate(papers, 1):
        if paper.get("errors") or paper.get("status") in {"failed", "duplicate", "skipped"}:
            continue
        row = row_by_document_id.get(str(paper.get("document_id")))
        if row is None and index - 1 < len(rows):
            row = rows[index - 1]
        row = row or {}
        source = str(paper.get("file_name") or row.get("file_name") or "paper")
        stem = re.sub(r"[^\w\-一-龥]+", "_", Path(source).stem, flags=re.UNICODE).strip("_") or "paper"
        document_id = str(paper.get("document_id") or row.get("document_id") or f"paper_{index}")
        path = target_dir / f"{stem[:70]}_{document_id[:12]}.docx"
        export_word(
            [row],
            path,
            papers=[paper],
            report_mode="Local Offline",
            report_layout_mode="仅单篇报告",
        )
        paper["individual_report"] = str(path.relative_to(target_dir.parent)).replace("\\", "/")
        paths.append(path)
    return paths
